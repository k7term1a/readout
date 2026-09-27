#!/usr/bin/env python3
"""Render a lab-meeting deck (deck.json) to an editable .pptx.

usage:
  python build_deck.py deck.json -o out.pptx
  python build_deck.py deck.json -o out.pptx --density visual --nav double
  python build_deck.py deck.json -o out.pptx --all-densities   # out-text/-balanced/-visual.pptx

The spec format is documented in references/spec.md.
"""
import argparse
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W, H = 13.333, 7.5
MARGIN = 0.9

DEFAULT_THEME = {
    "colors": {
        "active": "1085DE", "inactive": "B7DAF5", "inactive_double": "BFBFBF", "rule": "FF9900", "text": "595959", "dark": "222222",
        "problem": "E06666", "solution": "8CD96A", "neutral": "1085DE", "muted": "8A8A8A",
        "good_bg": "D9EAD3", "good_fg": "274E13", "bad_bg": "F4CCCC", "bad_fg": "990000",
        "tint": "EEF5FC", "ph_bg": "F3F6FA", "ph_line": "9FB3C8", "ph_text": "5A6B7D", "todo": "E69138",
    },
    "fonts": {"ea": "Microsoft JhengHei", "latin": "Arial", "ref_latin": "Times New Roman"},
}

DENSITIES = {
    # text_ratio: share of body width for the text column when a visual sits beside it
    "text":     {"text_ratio": 0.56, "size": 17, "min": 13, "key": "full", "auto_notes": False, "key_point": True},
    "balanced": {"text_ratio": 0.38, "size": 19, "min": 14, "key": "short", "auto_notes": True, "key_point": False},
    "visual":   {"text_ratio": 0.0,  "size": 18, "min": 16, "key": "one_line", "auto_notes": True, "key_point": False},
}
FALLBACK = {"full": ["full", "short", "one_line"], "short": ["short", "full", "one_line"],
            "one_line": ["one_line", "short", "full"]}

WARN = []


# ---------------------------------------------------------------- primitives
class Deck:
    def __init__(self, spec, base, density, nav_style):
        theme = json.loads(json.dumps(DEFAULT_THEME))
        for k, v in spec.get("theme", {}).items():
            theme.setdefault(k, {}).update(v)
        self.C, self.F = theme["colors"], theme["fonts"]
        user_colors = spec.get("theme", {}).get("colors", {})
        if "active" in user_colors and "neutral" not in user_colors:
            self.C["neutral"] = self.C["active"]
        self.spec, self.base = spec, base
        self.density, self.nav_style = density, nav_style
        self.chapters = spec.get("chapters") or []
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.body_top = 1.95 if nav_style == "single" else 2.2

    # colours / fonts
    def rgb(self, key):
        return RGBColor.from_string(self.C.get(key, key))

    def style(self, run, size, color, bold=False, latin=None):
        f = run.font
        f.size, f.bold, f.name = Pt(size), bold, latin or self.F["latin"]
        f.color.rgb = self.rgb(color)
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
            el.set("typeface", self.F["ea"])

    def slide(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def textbox(self, s, x, y, w, h, paras, size=18, color="text", bold=False, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP, bullet=False, space=10, latin=None, min_size=None, label=""):
        paras = paras if isinstance(paras, list) else [paras]
        paras = [p if isinstance(p, dict) else {"t": str(p)} for p in paras]
        if min_size:  # shrink to fit
            while size > min_size and est_height(paras, size, w, space, bullet) > h:
                size -= 1
            if est_height(paras, size, w, space, bullet) > h * 1.05:
                WARN.append(f"[{label}] text may overflow at {size}pt — shorten it or split the slide")
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(p.get("space", space))
            para.line_spacing = 1.15
            for seg, hl in split_marks(p["t"]):
                r = para.add_run()
                r.text = seg
                self.style(r, p.get("size", size), p.get("color", color), p.get("bold", bold) or hl, latin)
                if hl:
                    r.font.color.rgb = self.rgb("active")
            if bullet and p.get("bullet", True):
                ppr = para._p.get_or_add_pPr()
                ppr.set("marL", str(int(Inches(0.3))))
                ppr.set("indent", str(-int(Inches(0.3))))
                etree.SubElement(ppr, qn("a:buChar")).set("char", "•")
        return tb

    def shape(self, s, kind, x, y, w, h, fill=None, line=None, radius=None, dash=False):
        sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius is not None:
            sh.adjustments[0] = radius
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = self.rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = self.rgb(line)
            sh.line.width = Pt(1.25)
            if dash:
                sh.line.dash_style = MSO_LINE.DASH
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def pill(self, s, x, y, w, h, text, fill, size=16, color="FFFFFF", bold=True):
        sh = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, radius=0.5)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        self.style(r, size, color, bold)
        return sh

    def line(self, s, x1, y1, x2, y2, color="rule", weight=2.25, arrow=False):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = self.rgb(color)
        c.line.width = Pt(weight)
        if arrow:
            tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
            tail.set("type", "triangle")
        etree.SubElement(c._element.spPr, qn("a:effectLst"))  # no theme shadow
        return c

    # ------------------------------------------------------------ chrome
    def nav(self, s, chapter, section=None):
        tabs = [c["name"] if isinstance(c, dict) else c for c in self.chapters]
        if not tabs:
            return
        single = self.nav_style == "single"
        th, ty, size = (0.52, 0.24, 17) if single else (0.44, 0.16, 15)
        tw = min(1.75, (W - 1.2 - 0.3 * (len(tabs) - 1)) / len(tabs))
        tw = max(tw, max(text_w(t, size) + 0.4 for t in tabs)) if len(tabs) <= 6 else tw
        gap = 0.32
        total = len(tabs) * tw + (len(tabs) - 1) * gap
        if total > W - 1.2:
            gap = 0.15
            tw = (W - 1.2 - gap * (len(tabs) - 1)) / len(tabs)
            total = W - 1.2
        x = (W - total) / 2
        off = "inactive" if single else "inactive_double"
        for t in tabs:
            self.pill(s, x, ty, tw, th, t, "active" if t == chapter else off, size=size)
            x += tw + gap
        if single:
            self.line(s, 0.6, 0.98, W - 0.6, 0.98)
            return
        # double: section pills sit on the orange rule, which breaks around them
        rule_y, left, right = 0.95, W - 0.6, 0.6
        secs = self.sections_of(chapter)
        if secs:
            sh, ssize, sgap = 0.34, 13, 0.12
            widths = [text_w(t, ssize) + 0.45 for t in secs]
            x = left = 1.2  # left-aligned, leaving a short stub of rule before the first pill
            for t, w in zip(secs, widths):
                self.pill(s, x, rule_y - sh / 2, w, sh, t, "active" if t == section else off, size=ssize)
                x += w + sgap
            right = x - sgap
            left, right = left - 0.15, right + 0.15
        if secs:
            self.line(s, 0.6, rule_y, left, rule_y)
            self.line(s, right, rule_y, W - 0.6, rule_y)
        else:
            self.line(s, 0.6, rule_y, W - 0.6, rule_y)

    def sections_of(self, chapter):
        for c in self.chapters:
            if isinstance(c, dict) and c.get("name") == chapter:
                return c.get("sections", [])
        return []

    def heading(self, s, d, kind=None):
        """single: '/ coloured tag'; double: bold black title (+ optional grey subtitle)."""
        if self.nav_style == "single":
            self.tag(s, d.get("tag"), kind or d.get("kind", "neutral"))
            return
        if d.get("title"):
            self.textbox(s, 0.6, 1.25, W - 1.2, 0.5, d["title"], size=24, color="dark", bold=True,
                         anchor=MSO_ANCHOR.MIDDLE)
        if d.get("title") and d.get("subtitle"):
            self.textbox(s, 0.6, 1.75, W - 1.2, 0.3, d["subtitle"], size=14, color="muted")

    def page_number(self, s):
        """Bottom-right slide number (double nav only). Uses a slidenum field so it follows reordering."""
        if self.nav_style != "double":
            return
        tb = self.textbox(s, W - 1.1, H - 0.45, 0.7, 0.3, "", size=11, color="muted", align=PP_ALIGN.RIGHT)
        p = tb.text_frame.paragraphs[0]._p
        for r in p.findall(qn("a:r")):
            p.remove(r)
        fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
        rpr = etree.SubElement(fld, qn("a:rPr"), lang="zh-TW", sz="1100")
        fill = etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr"))
        fill.set("val", self.C["muted"])
        etree.SubElement(rpr, qn("a:latin"), typeface=self.F["latin"])
        etree.SubElement(rpr, qn("a:ea"), typeface=self.F["ea"])
        etree.SubElement(fld, qn("a:t")).text = str(len(self.prs.slides))
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            p.remove(end)
            p.append(end)

    def tag(self, s, text, kind):
        if not text:
            return
        y = self.body_top - 0.77
        self.textbox(s, 0.6, y, 0.3, 0.5, "/", size=18, color="dark", anchor=MSO_ANCHOR.MIDDLE)
        self.pill(s, 0.9, y, text_w(text, 17) + 0.7, 0.5, text, kind or "neutral", size=17)

    def body(self):
        return MARGIN, self.body_top, W - 2 * MARGIN, H - self.body_top - 0.5

    # ------------------------------------------------------------ visuals
    def figure(self, s, x, y, w, h, fig, label):
        fig = norm_fig(fig)
        path = fig.get("path")
        cap = f"{fig.get('id', '')}｜{fig.get('caption', '')}".strip("｜")
        if path:
            p = Path(path)
            p = p if p.is_absolute() else self.base / p
            if p.exists() and p.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp"):
                from PIL import Image
                ch = 0.35 if cap else 0
                with Image.open(p) as im:
                    ar = im.width / im.height
                iw, ih = w, w / ar
                if ih > h - ch:
                    ih = h - ch
                    iw = ih * ar
                ix, iy = x + (w - iw) / 2, y + (h - ch - ih) / 2
                s.shapes.add_picture(str(p), Inches(ix), Inches(iy), Inches(iw), Inches(ih))
                if cap:
                    self.textbox(s, x, iy + ih + 0.08, w, 0.3, cap, size=11, color="muted",
                                 align=PP_ALIGN.CENTER)
                return
            WARN.append(f"[{label}] image missing or unsupported (use PNG/JPG): {path} — placeholder drawn")
        sh = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill="ph_bg", line="ph_line",
                        radius=0.04, dash=True)
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, (t, sz, col, b) in enumerate([(cap or "Figure", 15, "ph_text", True),
                                             ("從論文 PDF 截圖後替換", 11, "muted", False)]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = t
            self.style(r, sz, col, b)

    def figures(self, s, x, y, w, h, figs, label, stack=False):
        n = len(figs)
        gap = 0.35
        if n == 0:
            return
        if stack and n <= 2:
            cells = [(x, y + k * ((h - gap * (n - 1)) / n + gap), w, (h - gap * (n - 1)) / n) for k in range(n)]
        elif n == 4 or (stack and n > 2):
            cw, ch = (w - gap) / 2, (h - gap) / 2
            cells = [(x + (k % 2) * (cw + gap), y + (k // 2) * (ch + gap), cw, ch) for k in range(min(n, 4))]
        else:
            n = min(n, 3)
            cw = (w - gap * (n - 1)) / n
            cells = [(x + k * (cw + gap), y, cw, h) for k in range(n)]
        if len(figs) > len(cells):
            WARN.append(f"[{label}] only {len(cells)} figures fit; move the rest to another slide")
        for f, cell in zip(figs, cells):
            self.figure(s, *cell, f, label)

    def table(self, s, x, y, w, rows, size=15):
        shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                 Inches(0.55 * len(rows)))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = shp.table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb("active" if r == 0 else ("FFFFFF" if r % 2 else "tint"))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
                rr = p.add_run()
                rr.text = str(val)
                col = "FFFFFF" if r == 0 else ("todo" if "待填" in str(val) else "dark")
                self.style(rr, size, col, bold=r == 0)

    def key_point(self, s, x, y, w, h, text):
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill="tint", radius=0.12)
        self.textbox(s, x + 0.3, y, w - 0.6, h, [{"t": "本頁重點", "size": 12, "color": "active", "bold": True},
                                                 {"t": text, "size": 17, "color": "dark", "bold": True}],
                     anchor=MSO_ANCHOR.MIDDLE, space=2)

    def notes(self, s, paras):
        if paras:
            s.notes_slide.notes_text_frame.text = "\n".join(p if isinstance(p, str) else str(p) for p in paras)

    # ------------------------------------------------------------ slide kinds
    def s_cover(self, d):
        style = d.get("style") or ("paper" if self.nav_style == "double" else "centered")
        if style not in ("centered", "paper"):
            WARN.append(f"[cover] unknown cover.style '{style}' — using 'centered'")
            style = "centered"
        s = self.slide()
        (self.cover_paper if style == "paper" else self.cover_centered)(s, d)
        self.notes(s, d.get("notes"))

    def cover_centered(self, s, d):
        """Single-nav cover: centred title, orange rule, then 作者．用途."""
        self.textbox(s, 1.2, 1.9, W - 2.4, 1.5, d.get("title", ""), size=38, color="dark", bold=True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, min_size=28, label="cover")
        self.line(s, 3.2, 3.6, W - 3.2, 3.6)
        y = 3.85
        for key, size, color in (("subtitle", 15, "muted"), ("byline", 17, "text"), ("date", 14, "muted")):
            if d.get(key):
                self.textbox(s, 1.2, y, W - 2.4, 0.5, d[key], size=size, color=color, align=PP_ALIGN.CENTER)
                y += 0.55

    def cover_paper(self, s, d):
        """Double-nav cover: venue badge, serif paper title, orange rule, authors | affiliations, presenter."""
        x, w = MARGIN, W - 2 * MARGIN
        if d.get("venue"):
            self.pill(s, x, 1.35, text_w(d["venue"], 14) + 0.5, 0.42, d["venue"], "active", size=14)
        self.textbox(s, x, 1.95, w, 1.65, d.get("paper_title") or d.get("title", ""), size=34, color="dark",
                     bold=True, anchor=MSO_ANCHOR.BOTTOM, latin=self.F["ref_latin"], min_size=24,
                     label="cover")
        self.line(s, x, 3.8, W - MARGIN, 3.8)
        # no author list (e.g. a progress report) -> fall back to the byline so the cover is not empty
        cols = [(as_list(d.get("authors") or d.get("byline")), "dark"), (as_list(d.get("affiliations")), "muted")]
        cw = (w - 0.5) / 2
        for k, (items, color) in enumerate(cols):
            if items:
                self.textbox(s, x + k * (cw + 0.5), 4.05, cw, 1.8, items, size=16, color=color, space=4,
                             latin=self.F["ref_latin"], min_size=12, label="cover")
        if d.get("presenter"):
            self.textbox(s, W - MARGIN - 5, H - 1.05, 5, 0.45, f"報告者：{d['presenter']}", size=17,
                         color="text", align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        if d.get("date"):
            self.textbox(s, x, H - 1.05, 4, 0.45, d["date"], size=14, color="muted", anchor=MSO_ANCHOR.MIDDLE)

    def s_content(self, d, label):
        dens = DENSITIES[d.get("density", self.density)]
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d)
        x, y, w, h = self.body()
        key = dens["key"]
        txt = pick_text(d, key)
        figs, table = d.get("figs", []), d.get("table")
        has_vis = bool(figs or table)
        is_list = txt is not None and txt is d.get("short") and isinstance(txt, list)

        if key == "one_line" or not txt:
            line_txt = pick_text(d, "one_line")
            if isinstance(line_txt, list):
                line_txt = line_txt[0]
            if not has_vis:  # statement slide
                self.textbox(s, x + 0.8, y, w - 1.6, h - 0.3, line_txt or "", size=30, color="dark", bold=True,
                             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=22, label=label)
            else:
                if line_txt:
                    self.textbox(s, x, y - 0.05, w, 0.5, line_txt, size=dens["size"], color="dark", bold=True,
                                 min_size=14, label=label)
                vy, vh = y + 0.65, h - 0.65
                if table and not figs:
                    self.table(s, x + 1.2, vy + 0.2, w - 2.4, table)
                else:
                    self.figures(s, x, vy, w, vh, figs, label)
        else:
            tw = w * (d.get("text_ratio") or dens["text_ratio"]) if has_vis else w
            kp = dens["key_point"] and d.get("one_line") and key == "full"
            th = h - 1.3 if kp else h
            self.textbox(s, x, y + 0.1, tw, th - 0.1, txt, size=dens["size"], bullet=is_list,
                         space=14 if is_list else 12, min_size=dens["min"], label=label)
            if kp:
                self.key_point(s, x, y + h - 1.0, tw, 1.0, d["one_line"])
            if has_vis:
                vx, vw = x + tw + 0.5, w - tw - 0.5
                if table:
                    self.table(s, vx, y + 0.2, vw, table, size=14 if vw < 6 else 15)
                    # figs beside a table are the visual-density alternative, not drawn here
                else:
                    self.figures(s, vx, y, vw, h, figs, label, stack=True)
        if d.get("callout"):
            self.textbox(s, x, H - 0.45, w, 0.3, d["callout"], size=12, color="muted")
        notes = d.get("notes") or (as_list(d.get("full")) if dens["auto_notes"] else [])
        self.notes(s, notes)

    def s_mapping(self, d, label):
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d, "neutral")
        pairs = d["pairs"]
        pw, ph = 3.6, 0.62
        lx, rx = 2.1, W - 2.1 - pw
        top = self.body_top + 0.1
        step = min(1.45, (H - top - 0.9) / max(len(pairs), 1))
        for k, (a, b) in enumerate(pairs):
            yy = top + k * step
            size = 17 if max(text_w(a, 17), text_w(b, 17)) < pw - 0.4 else 14
            self.pill(s, lx, yy, pw, ph, a, "problem", size=size)
            self.pill(s, rx, yy, pw, ph, b, "solution", size=size)
            self.line(s, lx + pw + 0.25, yy + ph / 2, rx - 0.25, yy + ph / 2, "444444", 1.75, arrow=True)
        self.notes(s, d.get("notes"))

    def s_split(self, d, label):
        dens = DENSITIES[d.get("density", self.density)]
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d, "neutral")
        x, y, w, _ = self.body()
        brief = dens["key"] == "one_line"
        h = 3.4 if brief else 3.9
        pw = (w - 0.8) / 2
        sides = [(d.get("good_title", "優點"), d.get("good", []), "good_bg", "good_fg", PP_ALIGN.LEFT),
                 (d.get("bad_title", "限制"), d.get("bad", []), "bad_bg", "bad_fg", PP_ALIGN.RIGHT)]
        for k, (title, items, bg, fg, align) in enumerate(sides):
            px = x + k * (pw + 0.8)
            self.shape(s, MSO_SHAPE.RECTANGLE, px, y, pw, h, fill=bg)
            self.textbox(s, px + 0.35, y + 0.3, pw - 0.7, 0.45, title, size=15, color="muted", bold=True,
                         align=align)
            paras = []
            for it in items:
                t, desc = (it + [""])[:2] if isinstance(it, list) else (it, "")
                paras.append({"t": t, "size": 19 if brief else 17, "color": fg, "bold": True})
                if desc and not brief:
                    paras.append({"t": desc, "size": 15, "color": "text"})
                    paras.append({"t": " ", "size": 8})
            self.textbox(s, px + 0.35, y + 0.95, pw - 0.7, h - 1.2, paras, align=align,
                         space=16 if brief else 6, min_size=12, size=17, label=label)
        self.line(s, W / 2, y + 0.2, W / 2, y + h - 0.2)
        auto = [f"{(it + [''])[0]}：{(it + [''])[1]}" for it in d.get("good", []) + d.get("bad", [])
                if isinstance(it, list)] if brief else []
        self.notes(s, d.get("notes") or auto)

    def s_refs(self, d, label):
        brief = DENSITIES[d.get("density", self.density)]["key"] == "one_line"
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d, "neutral")
        paras = []
        for r in d["refs"]:
            r = r if isinstance(r, dict) else {"title": r[0], "source": r[1] if len(r) > 1 else "",
                                                "gist": r[2] if len(r) > 2 else ""}
            paras.append({"t": f"–  {r['title']}", "size": 17, "color": "dark", "bold": True, "space": 3})
            if r.get("source"):
                paras.append({"t": f"    {r['source']}", "size": 13, "color": "muted", "space": 3})
            if r.get("gist") and not brief:
                paras.append({"t": f"    {r['gist']}", "size": 14, "color": "text", "space": 3})
            paras.append({"t": " ", "size": 8, "space": 3})
        top = self.body_top if d.get("title" if self.nav_style == "double" else "tag") else self.body_top - 0.45
        self.textbox(s, 1.2, top, W - 2.4, H - top - 0.4, paras, latin=self.F["ref_latin"], size=17,
                     min_size=11, label=label)
        self.notes(s, d.get("notes"))

    def s_statement(self, d, label):
        s = self.slide()
        if d.get("chapter"):
            self.nav(s, d.get("chapter"), d.get("section"))
        self.textbox(s, 1.5, 1.8, W - 3, 4.2, d.get("text", ""), size=40, color="dark", bold=True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=24, label=label)
        self.notes(s, d.get("notes"))

    def build(self):
        self.s_cover(self.spec.get("cover", {}))
        for i, d in enumerate(self.spec.get("slides", []), start=2):
            t = d.get("type", "content")
            label = f"slide {i} {d.get('tag') or t}"
            fn = getattr(self, f"s_{t}", None) if t != "cover" else None
            if fn is None:
                hint = " (divider pages were removed — delete this slide)" if t == "divider" else ""
                WARN.append(f"[{label}] unknown slide type '{t}' — skipped{hint}")
                continue
            fn(d, label)
            self.page_number(self.prs.slides[-1])
        return self.prs


# ---------------------------------------------------------------- helpers
def text_w(t, size):
    em = size / 72
    return sum(em if ord(ch) > 0x2E80 else em * 0.55 for ch in t)


def split_marks(t):
    """'a ==b== c' -> [('a ', False), ('b', True), (' c', False)]"""
    parts = t.split("==")
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p] or [("", False)]


def est_height(paras, size, width, space, bullet):
    total = 0
    for p in paras:
        s = p.get("size", size)
        avail = width - (0.3 if bullet else 0)
        lines = 0
        for seg in p["t"].replace("==", "").split("\n"):
            lines += max(1, -(-text_w(seg, s) // max(avail, 0.5)))
        total += lines * s * 1.15 / 72 + p.get("space", space) / 72
    return total


def as_list(v):
    if v is None:
        return []
    return v if isinstance(v, list) else [v]


def pick_text(d, key):
    for k in FALLBACK[key]:
        if d.get(k):
            return d[k]
    return None


def norm_fig(f):
    if isinstance(f, dict):
        return f
    f = list(f)
    return {"id": f[0], "caption": f[1] if len(f) > 1 else "", "path": f[2] if len(f) > 2 else None}


def render(spec, base, density, nav, out):
    WARN.clear()
    prs = Deck(spec, base, density, nav).build()
    prs.save(out)
    print(f"✔ {out}  ({len(prs.slides)} slides, density={density}, nav={nav})")
    for w in WARN:
        print("  ⚠", w)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("-o", "--out")
    ap.add_argument("--density", choices=list(DENSITIES))
    ap.add_argument("--nav", choices=["single", "double"])
    ap.add_argument("--all-densities", action="store_true")
    a = ap.parse_args()
    spec_path = Path(a.spec)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    style = spec.get("style", {})
    nav = a.nav or style.get("nav", "single")
    out = Path(a.out or spec_path.with_suffix(".pptx"))
    if a.all_densities:
        for d in DENSITIES:
            render(spec, spec_path.parent, d, nav, str(out.with_name(f"{out.stem}-{d}.pptx")))
    else:
        render(spec, spec_path.parent, a.density or style.get("density", "balanced"), nav, str(out))


if __name__ == "__main__":
    sys.exit(main())
