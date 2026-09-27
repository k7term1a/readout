"""Smoke tests: every example builds in every nav x density combination without layout warnings."""
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/lab-deck/scripts"))
import build_deck as bd  # noqa: E402

EXAMPLES = [ROOT / "skills/lab-deck/assets/example-deck.json", ROOT / "examples/progress-report/deck.json"]


@pytest.mark.parametrize("spec_path", EXAMPLES, ids=lambda p: p.parent.name)
@pytest.mark.parametrize("nav", ["single", "double"])
@pytest.mark.parametrize("density", list(bd.DENSITIES))
def test_builds_without_layout_warnings(tmp_path, spec_path, nav, density):
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out = tmp_path / "out.pptx"
    bd.render(spec, spec_path.parent, density, nav, str(out))
    layout_warnings = [w for w in bd.WARN if "image missing" not in w]
    assert not layout_warnings, layout_warnings
    assert len(Presentation(str(out)).slides) >= len([s for s in spec["slides"] if s.get("type") != "divider"])


def test_single_nav_skips_dividers(tmp_path):
    spec_path = EXAMPLES[0]
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out = tmp_path / "out.pptx"
    bd.render(spec, spec_path.parent, "balanced", "single", str(out))
    n_div = sum(s.get("type") == "divider" for s in spec["slides"])
    assert len(Presentation(str(out)).slides) == 1 + len(spec["slides"]) - n_div


def test_active_tab_matches_chapter(tmp_path):
    spec = {"chapters": ["甲", "乙"], "cover": {"title": "t"},
            "slides": [{"type": "content", "chapter": "乙", "tag": "x", "one_line": "y"}]}
    out = tmp_path / "out.pptx"
    bd.render(spec, tmp_path, "visual", "single", str(out))
    slide = Presentation(str(out)).slides[1]
    fills = {sh.text_frame.text: str(sh.fill.fore_color.rgb) for sh in slide.shapes
             if sh.has_text_frame and sh.text_frame.text in ("甲", "乙")}
    assert fills == {"乙": bd.DEFAULT_THEME["colors"]["active"], "甲": bd.DEFAULT_THEME["colors"]["inactive"]}
