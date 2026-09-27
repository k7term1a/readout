"""Smoke tests: every example builds in every nav x density combination without layout warnings."""
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.shapes.connector import Connector

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/2slide/scripts"))
import build_deck as bd  # noqa: E402

EXAMPLES = [ROOT / "skills/2slide/assets/example-deck.json", ROOT / "examples/progress-report/deck.json"]


@pytest.mark.parametrize("spec_path", EXAMPLES, ids=lambda p: p.parent.name)
@pytest.mark.parametrize("nav", ["single", "double"])
@pytest.mark.parametrize("density", list(bd.DENSITIES))
def test_builds_without_layout_warnings(tmp_path, spec_path, nav, density):
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out = tmp_path / "out.pptx"
    bd.render(spec, spec_path.parent, density, nav, str(out))
    layout_warnings = [w for w in bd.WARN if "image missing" not in w]
    assert not layout_warnings, layout_warnings
    expected = sum(1 + len(sl.get("focus", [])) if sl.get("focus") else 1 for sl in spec["slides"])
    assert len(Presentation(str(out)).slides) == 1 + expected  # cover + slides (+ focus pages)


def test_active_tab_matches_chapter(tmp_path):
    spec = {"chapters": ["甲", "乙"], "cover": {"title": "t"},
            "slides": [{"type": "content", "chapter": "乙", "tag": "x", "one_line": "y"}]}
    out = tmp_path / "out.pptx"
    bd.render(spec, tmp_path, "visual", "single", str(out))
    slide = Presentation(str(out)).slides[1]
    fills = {sh.text_frame.text: str(sh.fill.fore_color.rgb) for sh in slide.shapes
             if sh.has_text_frame and sh.text_frame.text in ("甲", "乙")}
    assert fills == {"乙": bd.DEFAULT_THEME["colors"]["active"], "甲": bd.DEFAULT_THEME["colors"]["inactive"]}


C = bd.DEFAULT_THEME["colors"]
DOUBLE_SPEC = {
    "chapters": [{"name": "介紹", "sections": ["背景", "動機"]}, {"name": "結論", "sections": []}],
    "cover": {"title": "t"},
    "slides": [{"type": "content", "chapter": "介紹", "section": "動機", "tag": "一、標籤", "kind": "problem",
                "title": "黑字標題", "subtitle": "灰色說明", "one_line": "y"},
               {"type": "content", "chapter": "結論", "tag": "x", "one_line": "z"}],
}


def build(tmp_path, spec, nav, density="visual"):
    out = tmp_path / "out.pptx"
    bd.render(spec, tmp_path, density, nav, str(out))
    return list(Presentation(str(out)).slides)


def texts(slide):
    return {sh.text_frame.text: sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text}


def test_double_nav_tabs_and_sections(tmp_path):
    t = texts(build(tmp_path, DOUBLE_SPEC, "double")[1])
    fill = {k: str(t[k].fill.fore_color.rgb) for k in ("介紹", "結論", "背景", "動機")}
    assert fill == {"介紹": C["active"], "結論": C["inactive_double"], "動機": C["active"], "背景": C["inactive_double"]}
    assert all(str(t[k].text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF" for k in fill)


def test_double_rule_breaks_around_sections(tmp_path):
    slides = build(tmp_path, DOUBLE_SPEC, "double")
    rules = lambda s: sorted((c.begin_x, c.end_x) for c in s.shapes if isinstance(c, Connector))
    with_secs, without = rules(slides[1]), rules(slides[2])
    pills = [sh for sh in slides[1].shapes if sh.has_text_frame and sh.text_frame.text in ("背景", "動機")]
    assert len(with_secs) == 2 and len(without) == 1
    (l0, l1), (r0, r1) = with_secs
    assert l1 < min(p.left for p in pills) and r0 > max(p.left + p.width for p in pills)
    # left-aligned but not flush: a short stub of rule sits before the first pill
    first = min(p.left for p in pills)
    assert l0 < l1 < first < bd.Inches(bd.W / 4)
    # the section pills straddle the rule
    rule_y = next(c for c in slides[1].shapes if isinstance(c, Connector)).begin_y
    assert all(p.top < rule_y < p.top + p.height for p in pills)


def test_double_uses_black_title_instead_of_tag(tmp_path):
    t = texts(build(tmp_path, DOUBLE_SPEC, "double")[1])
    assert "一、標籤" not in t and "/" not in t
    run = t["黑字標題"].text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == C["dark"]
    assert "灰色說明" in t
    single = texts(build(tmp_path, DOUBLE_SPEC, "single")[1])
    assert "一、標籤" in single and "黑字標題" not in single


def slidenum_fields(slide):
    return slide._element.findall(".//" + bd.qn("a:fld"))


def test_page_numbers_double_only_and_not_on_cover(tmp_path):
    slides = build(tmp_path, DOUBLE_SPEC, "double")
    assert not slidenum_fields(slides[0])
    assert [slidenum_fields(s)[0].get("type") for s in slides[1:]] == ["slidenum", "slidenum"]
    assert all(not slidenum_fields(s) for s in build(tmp_path, DOUBLE_SPEC, "single"))


def test_divider_is_rejected(tmp_path):
    spec = {**DOUBLE_SPEC, "slides": [{"type": "divider", "chapter": "介紹"}] + DOUBLE_SPEC["slides"]}
    assert len(build(tmp_path, spec, "double")) == 3
    assert any("divider" in w for w in bd.WARN)


def test_no_shadow_on_rules(tmp_path):
    for c in (sh for sh in build(tmp_path, DOUBLE_SPEC, "double")[1].shapes if isinstance(sh, Connector)):
        assert c._element.spPr.find(bd.qn("a:effectLst")) is not None


COVER = {"title": "中文標題", "byline": "報告人．用途", "venue": "NeurIPS 2022", "paper_title": "Paper Title",
         "authors": ["A. Author", "B. Author"], "affiliations": ["Some University"], "presenter": "王小明"}


def cover_of(tmp_path, nav, **cover):
    spec = {"chapters": ["甲"], "cover": {**COVER, **cover}, "slides": []}
    return texts(build(tmp_path, spec, nav)[0])


def test_cover_default_follows_nav(tmp_path):
    single = cover_of(tmp_path, "single")
    assert "中文標題" in single and "報告人．用途" in single and "NeurIPS 2022" not in single
    double = cover_of(tmp_path, "double")
    assert {"NeurIPS 2022", "Paper Title", "報告者：王小明"} <= set(double) and "中文標題" not in double


def test_cover_style_overrides_nav(tmp_path):
    assert "NeurIPS 2022" in cover_of(tmp_path, "single", style="paper")
    assert "中文標題" in cover_of(tmp_path, "double", style="centered")


def test_paper_cover_layout(tmp_path):
    t = cover_of(tmp_path, "double")
    badge, title = t["NeurIPS 2022"], t["Paper Title"]
    assert str(badge.fill.fore_color.rgb) == C["active"] and badge.top < title.top
    assert title.text_frame.paragraphs[0].runs[0].font.name == bd.DEFAULT_THEME["fonts"]["ref_latin"]
    authors, affil = t["A. Author\nB. Author"], t["Some University"]
    assert authors.top == affil.top and authors.left < affil.left  # two columns
    pres = t["報告者：王小明"]
    assert pres.left > bd.Inches(bd.W / 2) and pres.top > authors.top  # bottom-right
    assert "中文標題" in cover_of(tmp_path, "double", paper_title=None)


def test_unknown_cover_style_warns(tmp_path):
    cover_of(tmp_path, "double", style="fancy")
    assert any("cover.style" in w for w in bd.WARN)


def test_paper_cover_falls_back_to_byline(tmp_path):
    t = cover_of(tmp_path, "double", authors=None, paper_title=None)
    assert "中文標題" in t and "報告人．用途" in t


def bar_spec(conclusion, **extra):
    return {"chapters": ["甲"], "cover": {"title": "t"},
            "slides": [{"type": "content", "chapter": "甲", "tag": "x", "title": "標題", "full": ["段落"],
                        "short": ["條列"], "one_line": "一句話", "figs": [{"id": "Figure 1"}],
                        "conclusion": conclusion, **extra}]}


@pytest.mark.parametrize("nav", ["single", "double"])
@pytest.mark.parametrize("density", list(bd.DENSITIES))
def test_conclusion_bar_bottom_reserves_space(tmp_path, nav, density):
    slide = build(tmp_path, bar_spec("既有方法算不動", conclusion_pos="bottom"), nav, density)[1]
    bar = texts(slide)["既有方法算不動"]
    assert bar.name == "結論橫條" and str(bar.fill.fore_color.rgb) == C["problem"]
    assert str(bar.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert bar.left == bd.Inches(0.6) and bar.width == bd.Inches(bd.W - 1.2)  # full width, same as the rule
    # everything else (except the page number) ends above the bar
    others = [sh for sh in slide.shapes if sh.shape_id != bar.shape_id and not isinstance(sh, Connector)
              and not slidenum_fields_in(sh) and sh.top > bd.Inches(1.1)]
    assert others and all(sh.top + sh.height <= bar.top for sh in others)


def slidenum_fields_in(shape):
    return shape._element.findall(".//" + bd.qn("a:fld"))


def test_no_bar_without_conclusion(tmp_path):
    assert all(sh.name != "結論橫條" for sh in build(tmp_path, bar_spec(None), "double")[1].shapes)


def test_long_conclusion_warns(tmp_path):
    build(tmp_path, bar_spec("太長" * 60), "double")
    assert any("conclusion bar" in w for w in bd.WARN)


def named(slide, name):
    return [sh for sh in slide.shapes if sh.name == name]


def timing_targets(slide):
    """(spid, nodeType) of every entrance effect in the slide's timing tree."""
    out = []
    for ctn in slide._element.iter(bd.qn("p:cTn")):
        if ctn.get("presetClass") == "entr":
            spid = next(ctn.iter(bd.qn("p:spTgt"))).get("spid")
            out.append((int(spid), ctn.get("nodeType"), ctn.get("presetID")))
    return out


def test_conclusion_fades_in_on_one_click_with_veil(tmp_path):
    slide = build(tmp_path, bar_spec("結論"), "double", "balanced")[1]
    (veil,), (bar,) = named(slide, "結論遮罩"), named(slide, "結論橫條")
    # fade (presetID 10): veil on click, bar with it
    assert timing_targets(slide) == [(veil.shape_id, "clickEffect", "10"), (bar.shape_id, "withEffect", "10")]
    # translucent white veil from under the heading down to the bar, drawn above the content
    alpha = veil.fill._xPr.find(bd.qn("a:solidFill"))[0].find(bd.qn("a:alpha"))
    assert str(veil.fill.fore_color.rgb) == "FFFFFF" and int(alpha.get("val")) == bd.VEIL_ALPHA * 1000
    assert veil.left == bar.left and veil.width == bar.width
    title = texts(slide)["標題"]
    assert veil.top >= title.top + title.height
    # z-order: veil above every content shape, bar above the veil (only the page number may come later)
    order = [sh.shape_id for sh in slide.shapes if not slidenum_fields_in(sh)]
    assert order[-2:] == [veil.shape_id, bar.shape_id]


@pytest.mark.parametrize("kind", ["problem", "solution", "neutral"])
def test_conclusion_kind_colours(tmp_path, kind):
    (bar,) = named(build(tmp_path, bar_spec("結論", conclusion_kind=kind), "double")[1], "結論橫條")
    assert str(bar.fill.fore_color.rgb) == C[kind]


def test_keyframes_split_into_before_and_after(tmp_path):
    out = tmp_path / "kf.pptx"
    bd.render(bar_spec("結論"), tmp_path, "balanced", "double", str(out), keyframes=True)
    cover, before, after = Presentation(str(out)).slides
    assert not named(before, "結論橫條") and not named(before, "結論遮罩")
    assert named(after, "結論橫條") and named(after, "結論遮罩")
    assert not timing_targets(before) and not timing_targets(after)  # no animation in keyframe mode
    # the content does not move between frames
    geom = lambda sl: [(sh.left, sh.top, sh.width, sh.height) for sh in sl.shapes
                       if sh.name not in ("結論橫條", "結論遮罩") and not slidenum_fields_in(sh)]
    assert geom(before) == geom(after)


def near(a, b):
    return abs(a - b) <= 2  # EMU rounding


def content_geometry(slide):
    return [(sh.left, sh.top, sh.width, sh.height) for sh in slide.shapes
            if sh.name not in ("結論橫條", "結論遮罩") and not slidenum_fields_in(sh)]


@pytest.mark.parametrize("nav", ["single", "double"])
@pytest.mark.parametrize("density", list(bd.DENSITIES))
def test_conclusion_centered_in_veil_by_default(tmp_path, nav, density):
    slide = build(tmp_path, bar_spec("結論"), nav, density)[1]
    (veil,), (bar,) = named(slide, "結論遮罩"), named(slide, "結論橫條")
    assert abs((veil.top + veil.height / 2) - (bar.top + bar.height / 2)) <= 1  # vertically centred
    assert near(veil.top + veil.height, bd.Inches(bd.H - 0.5))  # veil covers the whole body
    # no space reserved: content is laid out exactly as without a conclusion
    plain = build(tmp_path, bar_spec(None), nav, density)[1]
    assert content_geometry(slide) == content_geometry(plain)


def test_conclusion_pos_bottom_veil_stops_at_bar(tmp_path):
    slide = build(tmp_path, bar_spec("結論", conclusion_pos="bottom"), "double")[1]
    (veil,), (bar,) = named(slide, "結論遮罩"), named(slide, "結論橫條")
    assert veil.top + veil.height == bar.top and bar.top + bar.height == bd.Inches(bd.H - 0.5)


def test_conclusion_pos_deck_default_and_override(tmp_path):
    spec = bar_spec("結論")
    spec["style"] = {"conclusion_pos": "bottom"}
    (bar,) = named(build(tmp_path, spec, "double")[1], "結論橫條")
    assert near(bar.top + bar.height, bd.Inches(bd.H - 0.5))
    spec["slides"][0]["conclusion_pos"] = "center"
    (veil,) = named(build(tmp_path, spec, "double")[1], "結論遮罩")
    assert near(veil.top + veil.height, bd.Inches(bd.H - 0.5))


def test_unknown_conclusion_pos_warns(tmp_path):
    build(tmp_path, bar_spec("結論", conclusion_pos="top"), "double")
    assert any("conclusion_pos" in w for w in bd.WARN)


CONCEPTS = {"圖像層級": 1, "區域層級": 2}


def concept_spec(slides, concepts=CONCEPTS):
    return {"chapters": ["甲"], "concepts": concepts, "cover": {"title": "t"}, "slides": slides}


def runs_of(slide):
    return [r for sh in slide.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs]


def test_parse_marks():
    assert bd.parse_marks("a ==b== {{c}} {{c|d}}") == [
        ("a ", False, None), ("b", True, None), (" ", False, None), ("c", False, "c"), (" ", False, None),
        ("d", False, "c")]
    assert bd.plain_text("{{圖像層級|整張圖}}的特徵") == "整張圖的特徵"


def test_inline_concept_text_is_coloured(tmp_path):
    spec = concept_spec([{"type": "content", "chapter": "甲", "tag": "x", "title": "{{區域層級}}標題",
                          "short": ["比較{{圖像層級}}與{{區域層級|區域}}"]}])
    slide = build(tmp_path, spec, "double", "balanced")[1]
    runs = {r.text: r for r in runs_of(slide)}
    assert "{{" not in "".join(runs)  # marks never leak into the slide
    assert str(runs["圖像層級"].font.color.rgb) == C["concept1_fg"] and runs["圖像層級"].font.bold
    assert str(runs["區域"].font.color.rgb) == C["concept2_fg"]
    assert str(runs["區域層級"].font.color.rgb) == C["concept2_fg"]  # in the title too


def test_figure_and_table_row_concept_blocks(tmp_path):
    spec = concept_spec([
        {"type": "content", "chapter": "甲", "one_line": "y", "figs": [{"id": "Figure 1", "concept": "區域層級"}]},
        {"type": "content", "chapter": "甲", "short": ["t"],
         "table": [["方法", "分數"], ["{{圖像層級}} A", "1"], ["B", "2"]]}])
    slides = build(tmp_path, spec, "double", "visual")
    ph = texts(slides[1])["Figure 1\n從論文 PDF 截圖後替換"]
    assert str(ph.fill.fore_color.rgb) == C["concept2_bg"] and str(ph.line.color.rgb) == C["concept2_fg"]
    slides = build(tmp_path, spec, "double", "balanced")
    tbl = next(sh for sh in slides[2].shapes if sh.has_table).table
    assert [str(tbl.cell(1, c).fill.fore_color.rgb) for c in range(2)] == [C["concept1_bg"]] * 2
    assert str(tbl.cell(2, 0).fill.fore_color.rgb) != C["concept1_bg"]
    assert tbl.cell(1, 0).text == "圖像層級 A"


def test_colour_table_lists_used_colours_and_pages(tmp_path):
    spec = concept_spec([
        {"type": "content", "chapter": "甲", "tag": "一、問題", "kind": "problem", "short": ["{{圖像層級}}"]},
        {"type": "mapping", "chapter": "甲", "pairs": [["問題 A", "解法 A"]]},
        {"type": "content", "chapter": "甲", "tag": "x", "short": ["{{圖像層級}} 再次出現"]},
        {"type": "content", "chapter": "甲", "tag": "x", "short": ["無"], "conclusion": "結論"}],
        concepts={**CONCEPTS, "沒用到": 3})
    rows = bd.render(spec, tmp_path, "balanced", "single", str(tmp_path / "o.pptx"))
    by = {r["color"]: r for r in rows}
    assert set(by) == {"紅", "綠", "概念色 1（藍）"}  # unused concept 2/3 not listed
    assert by["概念色 1（藍）"]["pages"] == [2, 4] and by["概念色 1（藍）"]["meaning"] == "圖像層級"
    assert by["概念色 1（藍）"]["hex"] == [C["concept1_bg"], C["concept1_fg"]]
    assert by["紅"]["meaning"] == "問題：問題 A" and by["紅"]["pages"] == [2, 3, 5]  # mapping names win
    md = bd.color_table_md(rows)
    assert "| 概念色 1（藍） | `B6CFF5`／`3C78D8` | 圖像層級 | 第 2、4 頁 |" in md
    assert bd.page_ranges([2, 3, 4, 7]) == "第 2–4、7 頁"


@pytest.mark.parametrize("concepts, needle", [({"a": 5}, "slot 5"), ({"a": 1, "b": 1}, "used twice")])
def test_bad_concept_slots_warn(tmp_path, concepts, needle):
    build(tmp_path, concept_spec([], concepts), "double")
    assert any(needle in w for w in bd.WARN)


def test_undeclared_concept_warns_and_renders_plain(tmp_path):
    slide = build(tmp_path, concept_spec([{"type": "content", "chapter": "甲", "short": ["{{未宣告}}"]}]),
                  "double", "balanced")[1]
    assert any("未宣告" in w for w in bd.WARN)
    assert "未宣告" in [r.text for r in runs_of(slide)]


def test_bullet_keeps_body_colour_before_concept(tmp_path):
    spec = concept_spec([{"type": "content", "chapter": "甲", "short": ["{{圖像層級}}：說明"]}])
    slide = build(tmp_path, spec, "double", "balanced")[1]
    clr = slide._element.find(".//" + bd.qn("a:buClr"))
    assert clr is not None and clr[0].get("val") == C["text"]


def test_concept_frame_around_real_image(tmp_path):
    from PIL import Image
    Image.new("RGB", (400, 300), "white").save(tmp_path / "f.png")
    spec = concept_spec([{"type": "content", "chapter": "甲", "one_line": "y",
                          "figs": [{"id": "Figure 1", "path": "f.png", "concept": "圖像層級"}]}])
    slide = build(tmp_path, spec, "double", "visual")[1]
    (frame,) = named(slide, "概念框/圖像層級")
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)  # picture
    assert str(frame.line.color.rgb) == C["concept1_fg"]
    assert frame.left < pic.left and frame.left + frame.width > pic.left + pic.width


# ---------------------------------------------------------------- screenshot marks
def mark_spec(tmp_path, marks, density="visual", **fig_extra):
    from PIL import Image
    Image.new("RGB", (1200, 640), "white").save(tmp_path / "t.png")
    fig = {"id": "Table 2", "path": "t.png", "marks": marks, **fig_extra}
    return concept_spec([{"type": "content", "chapter": "甲", "one_line": "y", "short": ["x"], "figs": [fig]}])


def mark_slide(tmp_path, marks, density="visual", nav="double", **fig_extra):
    return build(tmp_path, mark_spec(tmp_path, marks, **fig_extra), nav, density)[1]


def flat_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from flat_shapes(sh.shapes)
        else:
            yield sh


def by_name(slide):
    return {sh.name: sh for sh in flat_shapes(slide.shapes)}


def picture(slide):
    return next(sh for sh in flat_shapes(slide.shapes) if sh.shape_type == 13)


def test_mark_geometry_rows_and_cols():
    assert bd.mark_geometry({"type": "box", "row": 8, "rows": 8, "col": 3, "cols": 4}) == (0.5, 7 / 8, 0.25, 1 / 8)
    assert bd.mark_geometry({"type": "band", "row": [2, 5], "rows": 8}) == (0, 1 / 8, 1, 0.5)
    assert bd.mark_geometry({"type": "note", "x": 0.3, "y": 0.4}) == (0.3, 0.4, 0, 0)
    assert bd.mark_geometry({"type": "box", "x": 0.3}) is None


def test_box_sits_on_image_fractions(tmp_path):
    sl = mark_slide(tmp_path, [{"type": "box", "x": 0.5, "y": 0.25, "w": 0.2, "h": 0.1}])
    pic, box = picture(sl), by_name(sl)["標註/Table 2/紅框1"]
    assert near(box.left, pic.left + 0.5 * pic.width) and near(box.top, pic.top + 0.25 * pic.height)
    assert near(box.width, 0.2 * pic.width) and near(box.height, 0.1 * pic.height)
    assert str(box.line.color.rgb) == "FF0000" and box.fill.type == 5  # outline only (MSO_FILL_TYPE.BACKGROUND)


def test_band_is_translucent_concept_colour_with_outside_label(tmp_path):
    sl = mark_slide(tmp_path, [{"type": "band", "row": [2, 5], "rows": 8, "concept": "圖像層級", "text": "相同參數量"}])
    pic, n = picture(sl), by_name(sl)
    band, lab = n["標註/Table 2/色塊1"], n["標註/Table 2/色塊1說明"]
    assert str(band.fill.fore_color.rgb) == C["concept1_bg"]
    assert band.fill._xPr.find(bd.qn("a:solidFill"))[0].find(bd.qn("a:alpha")).get("val") == str(bd.BAND_ALPHA * 1000)
    assert near(band.width, pic.width)  # full width by default
    assert lab.left >= pic.left + pic.width  # label outside, to the right
    assert str(lab.text_frame.paragraphs[0].runs[0].font.color.rgb) == C["concept1_fg"]


@pytest.mark.parametrize("side", ["right", "left", "top", "bottom"])
def test_note_outside_with_arrow_to_point(tmp_path, side):
    sl = mark_slide(tmp_path, [{"type": "note", "x": 0.8, "y": 0.9, "text": "用執行效率換來的", "side": side}])
    pic, n = picture(sl), by_name(sl)
    txt, arrow = n["標註/Table 2/說明1"], n["標註/Table 2/說明1箭頭"]
    right, bottom = pic.left + pic.width, pic.top + pic.height
    assert {"right": txt.left >= right, "left": txt.left + txt.width <= pic.left,
            "top": txt.top + txt.height <= pic.top, "bottom": txt.top >= bottom}[side]
    assert near(arrow.end_x, pic.left + 0.8 * pic.width) and near(arrow.end_y, pic.top + 0.9 * pic.height)


def test_note_inside_sits_on_image(tmp_path):
    sl = mark_slide(tmp_path, [{"type": "note", "x": 0.8, "y": 0.9, "text": "說明", "side": "inside",
                                "tx": 0.4, "ty": 0.5}])
    pic, txt = picture(sl), by_name(sl)["標註/Table 2/說明1"]
    assert near(txt.left, pic.left + 0.4 * pic.width) and near(txt.top, pic.top + 0.5 * pic.height)


def test_static_marks_grouped_with_image(tmp_path):
    sl = mark_slide(tmp_path, [{"type": "box", "x": 0.1, "y": 0.1, "w": 0.1, "h": 0.1}])
    (grp,) = [sh for sh in sl.shapes if sh.shape_type == 6]
    assert grp.name == "標註/Table 2"
    assert {sh.shape_type for sh in grp.shapes} >= {13}  # the picture moved into the group
    assert not timing_targets(sl)


def test_reveal_click_steps_in_order(tmp_path):
    marks = [{"type": "band", "row": 2, "rows": 8, "concept": "圖像層級", "text": "A"},
             {"type": "box", "x": 0.5, "y": 0.8, "w": 0.2, "h": 0.1},
             {"type": "note", "x": 0.8, "y": 0.9, "text": "B", "with_previous": True}]
    sl = mark_slide(tmp_path, marks, reveal="click")
    assert not [sh for sh in sl.shapes if sh.shape_type == 6]  # animated marks are not grouped
    n = by_name(sl)
    order = [(spid, node) for spid, node, _ in timing_targets(sl)]
    ids = lambda *names: [n[f"標註/Table 2/{x}"].shape_id for x in names]
    assert [spid for spid, _ in order] == ids("色塊1", "色塊1說明", "紅框2", "說明3", "說明3箭頭")
    assert [node for _, node in order] == ["clickEffect", "withEffect", "clickEffect", "withEffect", "withEffect"]


def test_reveal_marks_then_conclusion_as_keyframes(tmp_path):
    spec = mark_spec(tmp_path, [{"type": "box", "x": 0.1, "y": 0.1, "w": 0.1, "h": 0.1},
                                {"type": "note", "x": 0.5, "y": 0.5, "text": "B"}], reveal="click")
    spec["slides"][0]["conclusion"] = "結論"
    out = tmp_path / "kf.pptx"
    bd.render(spec, tmp_path, "visual", "double", str(out), keyframes=True)
    frames = list(Presentation(str(out)).slides)[1:]
    has = lambda sl, name: any(sh.name == name for sh in flat_shapes(sl.shapes))
    assert len(frames) == 4  # nothing, +box, +note, +conclusion
    assert [has(f, "標註/Table 2/紅框1") for f in frames] == [False, True, True, True]
    assert [has(f, "標註/Table 2/說明2") for f in frames] == [False, False, True, True]
    assert [has(f, "結論橫條") for f in frames] == [False, False, False, True]


@pytest.mark.parametrize("mark, needle", [
    ({"type": "circle", "x": 0.1, "y": 0.1}, "needs type"),
    ({"type": "box", "x": 1.2, "y": 0.1, "w": 0.1, "h": 0.1}, "0–1 fractions"),
    ({"type": "note", "x": 0.1, "y": 0.1, "text": "a", "side": "middle"}, "side must be"),
])
def test_bad_marks_warn(tmp_path, mark, needle):
    mark_slide(tmp_path, [mark])
    assert any(needle in w for w in bd.WARN)


def test_small_annotated_figure_warns(tmp_path):
    mark_slide(tmp_path, [{"type": "note", "x": 0.1, "y": 0.1, "text": "說明"}], density="text")
    assert any("only" in w and "wide" in w for w in bd.WARN)


# ---------------------------------------------------------------- diagrams
DIAGRAM = {
    "direction": "right",
    "nodes": [{"id": "a", "label": "A", "col": 1, "row": 1}, {"id": "b", "label": "B", "col": 1, "row": 2},
              {"id": "c", "label": "C", "col": 2, "row": [1, 2]}, {"id": "d", "label": "D", "col": 3, "row": 1}],
    "edges": [["a", "c"], ["b", "c", "合併"], ["c", "d"]],
    "modules": [{"id": "m", "label": "模組", "nodes": ["a", "b"]}],
}


def diagram_slide(tmp_path, nav="double", diagram=None, **extra):
    spec = concept_spec([{"type": "diagram", "chapter": "甲", "tag": "t", "title": "架構", "diagram": "g", **extra}])
    spec["diagrams"] = {"g": diagram or DIAGRAM}
    return build(tmp_path, spec, nav)[1]


def emu_box(sh):
    return sh.left, sh.top, sh.width, sh.height


def test_diagram_is_one_named_group(tmp_path):
    sl = diagram_slide(tmp_path)
    (grp,) = [sh for sh in sl.shapes if sh.shape_type == 6]
    assert grp.name == "架構圖"
    n = by_name(sl)
    assert {"架構圖/a", "架構圖/b", "架構圖/c", "架構圖/d", "架構圖/模組/m", "架構圖/連線/a-c"} <= set(n)
    a, b, c, dd = (n[f"架構圖/{i}"] for i in "abcd")
    assert a.left == b.left < c.left < dd.left and a.top < b.top  # grid order
    assert near(c.top + c.height / 2, (a.top + b.top + b.height) / 2)  # spanning node centred on its rows
    assert n["架構圖/c"].text_frame.text == "C"  # label lives in the node shape


def cxn(sh):
    nv = sh._element.find(bd.qn("p:nvCxnSpPr")).find(bd.qn("p:cNvCxnSpPr"))
    st, en = nv.find(bd.qn("a:stCxn")), nv.find(bd.qn("a:endCxn"))
    return (int(st.get("id")), int(st.get("idx"))), (int(en.get("id")), int(en.get("idx")))


def connector_ends(sh):
    """Visual start/end (EMU) of a straight or bent connector, honouring rot=90° and flips."""
    x = sh._element.spPr.find(bd.qn("a:xfrm"))
    off, ext = x.find(bd.qn("a:off")), x.find(bd.qn("a:ext"))
    w, h = int(ext.get("cx")), int(ext.get("cy"))
    cx, cy = int(off.get("x")) + w / 2, int(off.get("y")) + h / 2
    su, sv = (w / 2 if x.get("flipH") == "1" else -w / 2), (h / 2 if x.get("flipV") == "1" else -h / 2)
    pts = [(su, sv), (-su, -sv)]
    if x.get("rot") == "5400000":
        pts = [(-v, u) for u, v in pts]
    return [(cx + u, cy + v) for u, v in pts]


def test_edges_glued_to_node_sides(tmp_path):
    n = by_name(diagram_slide(tmp_path))
    a, c = n["架構圖/a"], n["架構圖/c"]
    e = n["架構圖/連線/a-c"]
    assert cxn(e) == ((a.shape_id, 3), (c.shape_id, 1))  # right side -> left side
    (x1, y1), (x2, y2) = connector_ends(e)
    assert abs(x1 - (a.left + a.width)) < 2000 and abs(y1 - (a.top + a.height / 2)) < 2000
    assert abs(x2 - c.left) < 2000 and abs(y2 - (c.top + c.height / 2)) < 2000
    tail = e._element.spPr.find(bd.qn("a:ln")).find(bd.qn("a:tailEnd"))
    assert tail.get("type") == "triangle"


def test_down_direction_uses_vertical_first_elbows(tmp_path):
    g = {"direction": "down",
         "nodes": [{"id": "top", "label": "T", "col": [1, 3], "row": 1}, {"id": "l", "label": "L", "col": 1, "row": 2},
                   {"id": "r", "label": "R", "col": 3, "row": 2}],
         "edges": [["top", "l"], ["top", "r"]]}
    n = by_name(diagram_slide(tmp_path, diagram=g))
    top, left, right = n["架構圖/top"], n["架構圖/l"], n["架構圖/r"]
    for tgt, name in ((left, "top-l"), (right, "top-r")):
        e = n[f"架構圖/連線/{name}"]
        assert cxn(e) == ((top.shape_id, 2), (tgt.shape_id, 0))  # bottom -> top
        assert e._element.spPr.find(bd.qn("a:xfrm")).get("rot") == "5400000"
        (x1, y1), (x2, y2) = connector_ends(e)
        assert abs(y1 - (top.top + top.height)) < 2000 and abs(y2 - tgt.top) < 2000
        assert abs(x2 - (tgt.left + tgt.width / 2)) < 2000


def test_bend_clears_module_frame(tmp_path):
    n = by_name(diagram_slide(tmp_path))
    frame, e = n["架構圖/模組/m"], n["架構圖/連線/a-c"]
    (x1, _), (x2, _) = connector_ends(e)
    adj = e._element.spPr.find(bd.qn("a:prstGeom")).find(bd.qn("a:avLst"))[0].get("fmla")
    bend_x = x1 + int(adj.split()[1]) / 100000 * (x2 - x1)
    assert frame.left + frame.width < bend_x < x2  # turns after leaving the frame


def test_hide_keeps_layout_and_drops_edges(tmp_path):
    full = by_name(diagram_slide(tmp_path))
    part = by_name(diagram_slide(tmp_path, hide=["d"]))
    assert "架構圖/d" not in part and "架構圖/連線/c-d" not in part
    for i in "abc":
        assert near(part[f"架構圖/{i}"].left, full[f"架構圖/{i}"].left)
        assert near(part[f"架構圖/{i}"].top, full[f"架構圖/{i}"].top)


def test_highlight_relabel_and_colour_table(tmp_path):
    spec = concept_spec([{"type": "diagram", "chapter": "甲", "title": "t", "diagram": "g",
                          "highlight": ["c", "m"], "relabel": {"c": "新融合"}}])
    spec["diagrams"] = {"g": DIAGRAM}
    rows = bd.render(spec, tmp_path, "visual", "double", str(tmp_path / "o.pptx"))
    n = by_name(Presentation(str(tmp_path / "o.pptx")).slides[1])
    c = n["架構圖/c"]
    assert c.text_frame.text == "新融合"
    assert str(c.line.color.rgb) == C["new_line"] and str(c.fill.fore_color.rgb) == C["new_bg"]
    assert str(n["架構圖/模組/m"].line.color.rgb) == C["new_line"]
    green = next(r for r in rows if r["color"] == "綠色描邊")
    assert "新融合" in green["meaning"] and green["pages"] == [2]


def test_node_thumbnail_and_concept(tmp_path):
    from PIL import Image
    Image.new("RGB", (40, 30), "red").save(tmp_path / "x.png")
    g = {"nodes": [{"id": "i", "label": "輸入", "col": 1, "row": 1, "image": "x.png"},
                   {"id": "k", "label": "K", "col": 2, "row": 1, "concept": "圖像層級"}], "edges": [["i", "k"]]}
    n = by_name(diagram_slide(tmp_path, diagram=g))
    assert n["架構圖/i/縮圖"].shape_type == 13 and n["架構圖/i/文字"].text_frame.text == "輸入"
    assert str(n["架構圖/k"].fill.fore_color.rgb) == C["concept1_bg"]


@pytest.mark.parametrize("extra, diagram, needle", [
    ({"diagram": "nope"}, None, "not defined"),
    ({"highlight": ["zz"]}, None, "no node or module 'zz'"),
    ({}, {"nodes": [{"id": "a", "col": 1, "row": 1}], "edges": [["a", "q"]]}, "unknown node"),
    ({}, {"nodes": [{"id": str(i), "col": i, "row": 1} for i in range(1, 12)]}, "too dense"),
])
def test_diagram_warnings(tmp_path, extra, diagram, needle):
    spec = concept_spec([{"type": "diagram", "chapter": "甲", "title": "t", "diagram": "g", **extra}])
    spec["diagrams"] = {"g": diagram or DIAGRAM}
    build(tmp_path, spec, "double")
    assert any(needle in w for w in bd.WARN), bd.WARN


def test_diagram_with_conclusion_keyframes(tmp_path):
    spec = concept_spec([{"type": "diagram", "chapter": "甲", "title": "t", "diagram": "g", "conclusion": "晚期融合"}])
    spec["diagrams"] = {"g": DIAGRAM}
    out = tmp_path / "kf.pptx"
    bd.render(spec, tmp_path, "visual", "double", str(out), keyframes=True)
    before, after = list(Presentation(str(out)).slides)[1:]
    assert not named(before, "結論橫條") and named(after, "結論橫條")


# ---------------------------------------------------------------- timeline
TL = {"type": "timeline", "chapter": "甲", "tag": "t", "title": "演進",
      "lanes": [{"name": "圖像層級", "concept": "圖像層級"}, {"name": "其他"}],
      "items": [{"label": "CLIP", "year": 2021, "lane": "圖像層級", "note": "ICML"},
                {"label": "ALIGN", "year": 2021, "lane": "圖像層級"},
                {"label": "BLIP", "year": 2022, "lane": "圖像層級"},
                {"label": "X", "year": 2020, "lane": "其他"},
                {"label": "本文", "year": 2023, "lane": "其他", "highlight": True}]}


def timeline_slide(tmp_path, nav="double", **extra):
    return build(tmp_path, concept_spec([{**TL, **extra}]), nav)[1]


def centre_x(sh):
    return sh.left + sh.width / 2


def test_timeline_structure_and_colours(tmp_path):
    sl = timeline_slide(tmp_path)
    (grp,) = [sh for sh in sl.shapes if sh.shape_type == 6]
    assert grp.name == "時間軸"
    n = by_name(sl)
    assert str(n["時間軸/分類/圖像層級"].fill.fore_color.rgb) == C["concept1_bg"]
    assert str(n["時間軸/分類/其他"].fill.fore_color.rgb) == C["lane_bg"]  # no concept -> grey, not a new colour
    assert str(n["時間軸/本文"].line.color.rgb) == C["new_line"]
    assert n["時間軸/CLIP/註記"].text_frame.text == "ICML"


def test_timeline_items_sit_on_their_year(tmp_path):
    n = by_name(timeline_slide(tmp_path))
    ticks = {yr: n[f"時間軸/年份軸/{yr}"] for yr in range(2020, 2024)}
    for label, yr in (("CLIP", 2021), ("BLIP", 2022), ("X", 2020), ("本文", 2023)):
        assert abs(centre_x(n[f"時間軸/{label}"]) - ticks[yr].begin_x) < 2000
    # same year, same band: stacked, not overlapping
    a, c = n["時間軸/ALIGN"], n["時間軸/CLIP"]
    assert a.top + a.height <= c.top or c.top + c.height <= a.top
    # items stay inside their band
    band = n["時間軸/分類/圖像層級"]
    for lab in ("CLIP", "ALIGN", "BLIP"):
        sh = n[f"時間軸/{lab}"]
        assert band.top <= sh.top and sh.top + sh.height <= band.top + band.height


def test_timeline_years_range_and_colour_table(tmp_path):
    spec = concept_spec([{**TL, "years": [2018, 2024]}])
    rows = bd.render(spec, tmp_path, "visual", "double", str(tmp_path / "o.pptx"))
    n = by_name(Presentation(str(tmp_path / "o.pptx")).slides[1])
    assert "時間軸/年份軸/2018" in n and "時間軸/年份軸/2024" in n
    green = next(r for r in rows if r["color"] == "綠色描邊")
    assert "本文" in green["meaning"]


@pytest.mark.parametrize("items, needle", [
    ([{"label": "A", "year": 2021.5}], "integer year"),
    ([{"label": "A", "year": 2021, "lane": "不存在"}], "unknown lane"),
    ([{"label": f"模型{i}", "year": 2021, "lane": "其他"} for i in range(12)], "too crowded"),
])
def test_timeline_warnings(tmp_path, items, needle):
    timeline_slide(tmp_path, items=items)
    assert any(needle in w for w in bd.WARN), bd.WARN


def test_timeline_conclusion_keyframes(tmp_path):
    out = tmp_path / "kf.pptx"
    bd.render(concept_spec([{**TL, "conclusion": "都只做圖像層級"}]), tmp_path, "visual", "double", str(out),
              keyframes=True)
    before, after = list(Presentation(str(out)).slides)[1:]
    assert not named(before, "結論橫條") and named(after, "結論橫條")


# ---------------------------------------------------------------- progressive focus
def alpha_of(sh, path=("a:solidFill", "a:srgbClr")):
    clr = sh._element.find(bd.qn("p:spPr")).find("/".join(bd.qn(x) for x in path))
    a = clr.find(bd.qn("a:alpha")) if clr is not None else None
    return int(a.get("val")) if a is not None else 100000


def faded(sh):
    runs = [r for p in sh.text_frame.paragraphs for r in p.runs] if sh.has_text_frame else []
    return alpha_of(sh) < 100000 or any(str(r.font.color.rgb) == bd.FADE_TEXT for r in runs)


def focus_deck(tmp_path, focus, **extra):
    spec = concept_spec([{"type": "diagram", "chapter": "甲", "title": "元件", "diagram": "g", "focus": focus, **extra}])
    spec["diagrams"] = {"g": DIAGRAM}
    out = tmp_path / "f.pptx"
    bd.render(spec, tmp_path, "visual", "double", str(out))
    return list(Presentation(str(out)).slides)[1:]


def test_focus_expands_overview_plus_one_page_per_step(tmp_path):
    pages = focus_deck(tmp_path, [{"on": "c", "subtitle": "先看 C"}, {"on": "d"}])
    assert len(pages) == 3
    overview, on_c, on_d = (by_name(p) for p in pages)
    assert not any(faded(sh) for name, sh in overview.items() if name.startswith("架構圖/"))
    assert not faded(on_c["架構圖/c"]) and faded(on_c["架構圖/a"]) and faded(on_c["架構圖/d"])
    assert "先看 C" in texts(pages[1])
    # edges touching the lit node stay; others fade
    assert alpha_of(on_c["架構圖/連線/a-c"], ("a:ln", "a:solidFill", "a:srgbClr")) == 100000
    assert not faded(on_d["架構圖/d"]) and faded(on_d["架構圖/c"])
    # positions never move between pages
    assert emu_box(on_c["架構圖/c"]) == emu_box(on_d["架構圖/c"]) == emu_box(overview["架構圖/c"])


def test_focus_on_module_lights_its_nodes(tmp_path):
    (_, page) = focus_deck(tmp_path, [{"on": "m"}])
    n = by_name(page)
    assert not faded(n["架構圖/a"]) and not faded(n["架構圖/b"]) and not faded(n["架構圖/模組/m"])
    assert faded(n["架構圖/c"])


def test_focus_overview_can_be_skipped_and_conclusion_not_inherited(tmp_path):
    pages = focus_deck(tmp_path, [{"on": "a"}, {"on": "b", "conclusion": "B 才是關鍵"}],
                       focus_overview=False, conclusion="總結")
    assert len(pages) == 2
    assert not named(pages[0], "結論橫條")  # the slide-level conclusion is not copied onto steps
    assert texts(pages[1])["B 才是關鍵"]


def test_focus_one_block_per_step_and_unknown_id(tmp_path):
    focus_deck(tmp_path, [{"on": ["a", "b"]}, {"on": "zz"}])
    assert any("one block per step" in w for w in bd.WARN)
    assert any("'zz' matches no" in w for w in bd.WARN)


def test_focus_on_screenshot_region(tmp_path):
    from PIL import Image
    Image.new("RGB", (900, 300), "white").save(tmp_path / "f.png")
    fig = {"id": "Figure 3", "path": "f.png",
           "regions": [{"id": "a", "x": 0, "y": 0, "w": 1 / 3, "h": 1}, {"id": "b", "x": 1 / 3, "y": 0, "w": 1 / 3, "h": 1}],
           "marks": [{"type": "box", "x": 0.8, "y": 0.2, "w": 0.1, "h": 0.2},
                     {"type": "box", "x": 0.1, "y": 0.2, "w": 0.1, "h": 0.2}]}
    spec = concept_spec([{"type": "content", "chapter": "甲", "title": "t", "one_line": "y", "figs": [fig],
                          "focus": [{"on": "a"}, {"on": "b"}]}])
    out = tmp_path / "f.pptx"
    bd.render(spec, tmp_path, "visual", "double", str(out))
    overview, on_a, on_b = list(Presentation(str(out)).slides)[1:]
    assert not [sh for sh in flat_shapes(overview.shapes) if sh.name.startswith("聚焦/")]
    n = by_name(on_a)
    pic = picture(on_a)
    veils = [sh for name, sh in n.items() if name.startswith("聚焦/Figure 3/遮罩")]
    assert veils and all(alpha_of(v) == (100 - bd.FADE_ALPHA) * 1000 for v in veils)
    # the lit third (left) is not covered by any veil
    lit_right = pic.left + pic.width / 3
    assert all(v.left >= lit_right - 2000 or v.top >= pic.top + pic.height - 2000 for v in veils)
    # a mark inside region a stays; one outside fades
    assert not faded(n["標註/Figure 3/紅框2"]) and alpha_of(n["標註/Figure 3/紅框1"], ("a:ln", "a:solidFill", "a:srgbClr")) < 100000


def test_more_than_three_clicks_warns(tmp_path):
    marks = [{"type": "box", "x": 0.1 * k, "y": 0.1, "w": 0.05, "h": 0.05} for k in range(1, 5)]
    mark_slide(tmp_path, marks, reveal="click")
    assert any("click animations on one slide" in w for w in bd.WARN)


def test_single_nav_shows_focus_subtitle_next_to_tag(tmp_path):
    spec = concept_spec([{"type": "diagram", "chapter": "甲", "tag": "元件", "title": "t", "diagram": "g",
                          "focus": [{"on": "c", "subtitle": "先看 C"}]}])
    spec["diagrams"] = {"g": DIAGRAM}
    step = build(tmp_path, spec, "single")[2]
    t = texts(step)
    assert "先看 C" in t and t["先看 C"].left > t["元件"].left + t["元件"].width
    overview = texts(build(tmp_path, spec, "single")[1])
    assert "先看 C" not in overview  # ordinary single-nav pages keep the tag-only heading
