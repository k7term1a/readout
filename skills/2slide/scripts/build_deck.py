#!/usr/bin/env python3
"""Render a lab-meeting deck (deck.json) to an editable .pptx.

usage:
  python build_deck.py deck.json -o out.pptx
  python build_deck.py deck.json -o out.pptx --density visual --nav double
  python build_deck.py deck.json -o out.pptx --all-densities   # out-text/-balanced/-visual.pptx
  python build_deck.py deck.json -o out.pptx --keyframes       # conclusion bars as before/after slides

The spec format is documented in references/spec.md.
"""
import argparse
import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W, H = 13.333, 7.5
MARGIN = 0.9
BAR_H = 0.62  # conclusion bar
VEIL_ALPHA = 70  # % opacity of the white veil that fades the content when the conclusion bar appears
BAND_ALPHA = 45  # % opacity of a concept band laid over a screenshot
FADE_ALPHA = 30  # % opacity left on parts that are out of focus (style.md: ~70% transparent)
FADE_TEXT = "C0C0C0"
MAX_MARKS = 3  # annotation groups (red box, band, note; appear+disappear counts once) per slide
FOCUS_KEYS = ("title", "subtitle", "tag", "kind", "one_line", "short", "full", "notes",
              "conclusion", "conclusion_kind", "conclusion_pos")
NOTE_SIDES = ("right", "left", "top", "bottom", "inside")

DEFAULT_THEME = {
    "colors": {
        "active": "1085DE", "inactive": "B7DAF5", "inactive_double": "BFBFBF", "rule": "FF9900", "text": "595959", "dark": "222222",
        "problem": "E06666", "solution": "8CD96A", "neutral": "1085DE", "muted": "8A8A8A",
        "good_bg": "D9EAD3", "good_fg": "274E13", "bad_bg": "F4CCCC", "bad_fg": "990000",
        "tint": "EEF5FC", "ph_bg": "F3F6FA", "ph_line": "9FB3C8", "ph_text": "5A6B7D", "todo": "E69138",
        "mark": "FF0000", "new_line": "6AA84F", "new_bg": "D9EAD3", "node_line": "7F7F7F", "edge": "595959",
        "lane_bg": "EFEFEF", "lane_fg": "7F7F7F",
        "concept1_bg": "B6CFF5", "concept1_fg": "3C78D8", "concept2_bg": "FFECB3", "concept2_fg": "BF9000",
        "concept3_bg": "D9D2E9", "concept3_fg": "674EA7", "concept4_bg": "D0E0E3", "concept4_fg": "45818E",
    },
    "fonts": {"ea": "Microsoft JhengHei", "latin": "Arial", "ref_latin": "Times New Roman"},
}

DENSITIES = {
    # text_ratio: share of body width for the text column when a visual sits beside it
    "text":     {"text_ratio": 0.56, "size": 17, "min": 13, "key": "full", "auto_notes": False, "key_point": True},
    "balanced": {"text_ratio": 0.38, "size": 19, "min": 14, "key": "short", "auto_notes": True, "key_point": False},
    "visual":   {"text_ratio": 0.0,  "size": 18, "min": 16, "key": "one_line", "auto_notes": True, "key_point": False},
}
FALLBACK = {"full": ["full", "short", "one_line"], "short": ["short", "full", "one_line"],
            "one_line": ["one_line", "short", "full"]}

CONCEPT_NAMES = {1: "藍", 2: "黃", 3: "紫", 4: "青灰"}

WARN = []


# ---------------------------------------------------------------- primitives
class Deck:
    def __init__(self, spec, base, density, nav_style, keyframes=False):
        theme = json.loads(json.dumps(DEFAULT_THEME))
        for k, v in spec.get("theme", {}).items():
            theme.setdefault(k, {}).update(v)
        self.C, self.F = theme["colors"], theme["fonts"]
        user_colors = spec.get("theme", {}).get("colors", {})
        if "active" in user_colors and "neutral" not in user_colors:
            self.C["neutral"] = self.C["active"]
        self.spec, self.base = spec, base
        self.density, self.nav_style = density, nav_style
        self.keyframes = keyframes
        self.chapters = spec.get("chapters") or []
        self.concepts = {}  # name -> slot 1..4
        for name, slot in (spec.get("concepts") or {}).items():
            if slot not in CONCEPT_NAMES:
                WARN.append(f"[concepts] '{name}' has slot {slot!r}; use 1–4 — ignored")
            elif slot in self.concepts.values():
                WARN.append(f"[concepts] slot {slot} is used twice ('{name}') — each concept needs its own colour")
            else:
                self.concepts[name] = slot
        # colour key -> {"pages": set, "names": [...], "mapping": [...]} for the colour table
        self.uses = {}
        self.page = 0
        self.steps = []
        self.focus = None  # id lit on the current focus page (None = nothing faded)
        self.focus_hit = False
        self.in_focus_seq = False
        self.pending_out = []
        self.mark_units = 0
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.body_top = 1.95 if nav_style == "single" else 2.2

    # colours / fonts
    def rgb(self, key):
        return RGBColor.from_string(self.C.get(key, key))

    def style(self, run, size, color, bold=False, latin=None):
        f = run.font
        f.size, f.bold, f.name = Pt(size), bold, latin or self.F["latin"]
        f.color.rgb = self.rgb(color)
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
            el.set("typeface", self.F["ea"])

    def slide(self):
        sl = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.page = len(self.prs.slides)
        self.steps = []  # click steps on this slide: [{"in": [ids], "out": [ids]}, ...]
        self.pending_out = []  # shapes that fade out on the next click (marks with "exit": true)
        self.mark_units = 0
        return sl

    def add_step(self, ids, merge=False):
        """One click: fade in `ids` (and fade out whatever an earlier mark asked to leave on this click)."""
        if merge and self.steps:
            self.steps[-1]["in"].extend(ids)
            return
        self.steps.append({"in": list(ids), "out": self.pending_out})
        self.pending_out = []

    # ------------------------------------------------------------ marks & colour usage
    def use(self, key, name=None, source="names"):
        u = self.uses.setdefault(key, {"pages": set(), "names": [], "mapping": []})
        u["pages"].add(self.page)
        if name and name not in u[source]:
            u[source].append(name)

    def segments(self, t):
        """Parse ==highlight== and {{concept}} / {{concept|text}} marks; logs concept use on this page."""
        out = []
        for text, hl, concept in parse_marks(str(t)):
            if concept is not None:
                if concept not in self.concepts:
                    msg = f"[concepts] '{concept}' is not declared in top-level concepts — drawn as plain text"
                    if msg not in WARN:
                        WARN.append(msg)
                    concept = None
                else:
                    self.use(f"concept:{concept}", concept)
            out.append((text, hl, concept))
        return out

    def plain(self, t):
        return "".join(seg for seg, _, _ in self.segments(t))

    def concept_slot(self, name, label):
        if not name:
            return None
        if name not in self.concepts:
            WARN.append(f"[{label}] concept '{name}' is not declared in top-level concepts — ignored")
            return None
        self.use(f"concept:{name}", name)
        return self.concepts[name]

    def textbox(self, s, x, y, w, h, paras, size=18, color="text", bold=False, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP, bullet=False, space=10, latin=None, min_size=None, label=""):
        paras = paras if isinstance(paras, list) else [paras]
        paras = [p if isinstance(p, dict) else {"t": str(p)} for p in paras]
        if min_size:  # shrink to fit
            while size > min_size and est_height(paras, size, w, space, bullet) > h:
                size -= 1
            if est_height(paras, size, w, space, bullet) > h * 1.05:
                WARN.append(f"[{label}] text may overflow at {size}pt — shorten it or split the slide")
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(p.get("space", space))
            para.line_spacing = 1.15
            for seg, hl, concept in self.segments(p["t"]):
                r = para.add_run()
                r.text = seg
                self.style(r, p.get("size", size), p.get("color", color),
                           p.get("bold", bold) or hl or concept is not None, latin)
                if hl:
                    r.font.color.rgb = self.rgb("active")
                if concept is not None:
                    r.font.color.rgb = self.rgb(f"concept{self.concepts[concept]}_fg")
            if bullet and p.get("bullet", True):
                ppr = para._p.get_or_add_pPr()
                ppr.set("marL", str(int(Inches(0.3))))
                ppr.set("indent", str(-int(Inches(0.3))))
                # bullet keeps the body colour even when the line starts with a coloured concept
                bu_clr = etree.SubElement(etree.SubElement(ppr, qn("a:buClr")), qn("a:srgbClr"))
                bu_clr.set("val", self.C.get(p.get("color", color), p.get("color", color)))
                etree.SubElement(ppr, qn("a:buChar")).set("char", "•")
        return tb

    def shape(self, s, kind, x, y, w, h, fill=None, line=None, radius=None, dash=False):
        sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius is not None:
            sh.adjustments[0] = radius
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = self.rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = self.rgb(line)
            sh.line.width = Pt(1.25)
            if dash:
                sh.line.dash_style = MSO_LINE.DASH
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def pill(self, s, x, y, w, h, text, fill, size=16, color="FFFFFF", bold=True, source="names"):
        text = self.plain(text)
        if fill in ("problem", "solution"):
            self.use(fill, text, source)
        sh = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, radius=0.5)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        self.style(r, size, color, bold)
        return sh

    def line(self, s, x1, y1, x2, y2, color="rule", weight=2.25, arrow=False):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = self.rgb(color)
        c.line.width = Pt(weight)
        if arrow:
            tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
            tail.set("type", "triangle")
        etree.SubElement(c._element.spPr, qn("a:effectLst"))  # no theme shadow
        return c

    # ------------------------------------------------------------ chrome
    def nav(self, s, chapter, section=None):
        tabs = [c["name"] if isinstance(c, dict) else c for c in self.chapters]
        if not tabs:
            return
        single = self.nav_style == "single"
        th, ty, size = (0.52, 0.24, 17) if single else (0.44, 0.16, 15)
        tw = min(1.75, (W - 1.2 - 0.3 * (len(tabs) - 1)) / len(tabs))
        tw = max(tw, max(text_w(t, size) + 0.4 for t in tabs)) if len(tabs) <= 6 else tw
        gap = 0.32
        total = len(tabs) * tw + (len(tabs) - 1) * gap
        if total > W - 1.2:
            gap = 0.15
            tw = (W - 1.2 - gap * (len(tabs) - 1)) / len(tabs)
            total = W - 1.2
        x = (W - total) / 2
        off = "inactive" if single else "inactive_double"
        for t in tabs:
            self.pill(s, x, ty, tw, th, t, "active" if t == chapter else off, size=size)
            x += tw + gap
        if single:
            self.line(s, 0.6, 0.98, W - 0.6, 0.98)
            return
        # double: section pills sit on the orange rule, which breaks around them
        rule_y, left, right = 0.95, W - 0.6, 0.6
        secs = self.sections_of(chapter)
        if secs:
            sh, ssize, sgap = 0.34, 13, 0.12
            widths = [text_w(t, ssize) + 0.45 for t in secs]
            x = left = 1.2  # left-aligned, leaving a short stub of rule before the first pill
            for t, w in zip(secs, widths):
                self.pill(s, x, rule_y - sh / 2, w, sh, t, "active" if t == section else off, size=ssize)
                x += w + sgap
            right = x - sgap
            left, right = left - 0.15, right + 0.15
        if secs:
            self.line(s, 0.6, rule_y, left, rule_y)
            self.line(s, right, rule_y, W - 0.6, rule_y)
        else:
            self.line(s, 0.6, rule_y, W - 0.6, rule_y)

    def sections_of(self, chapter):
        for c in self.chapters:
            if isinstance(c, dict) and c.get("name") == chapter:
                return c.get("sections", [])
        return []

    def heading(self, s, d, kind=None):
        """single: '/ coloured tag'; double: bold black title (+ optional grey subtitle)."""
        if self.nav_style == "single":
            self.tag(s, d.get("tag"), kind or d.get("kind", "neutral"))
            if d.get("_focus") and d.get("tag") and d.get("subtitle"):  # a focus step's explanation
                x = 0.9 + text_w(d["tag"], 17) + 0.7 + 0.3
                self.textbox(s, x, self.body_top - 0.77, W - 0.6 - x, 0.5, d["subtitle"], size=14, color="muted",
                             anchor=MSO_ANCHOR.MIDDLE)
            return
        if d.get("title"):
            self.textbox(s, 0.6, 1.25, W - 1.2, 0.5, d["title"], size=24, color="dark", bold=True,
                         anchor=MSO_ANCHOR.MIDDLE)
        if d.get("title") and d.get("subtitle"):
            self.textbox(s, 0.6, 1.75, W - 1.2, 0.3, d["subtitle"], size=14, color="muted")

    def page_number(self, s):
        """Bottom-right slide number (double nav only). Uses a slidenum field so it follows reordering."""
        if self.nav_style != "double":
            return
        tb = self.textbox(s, W - 1.1, H - 0.45, 0.7, 0.3, "", size=11, color="muted", align=PP_ALIGN.RIGHT)
        p = tb.text_frame.paragraphs[0]._p
        for r in p.findall(qn("a:r")):
            p.remove(r)
        fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
        rpr = etree.SubElement(fld, qn("a:rPr"), lang="zh-TW", sz="1100")
        fill = etree.SubElement(etree.SubElement(rpr, qn("a:solidFill")), qn("a:srgbClr"))
        fill.set("val", self.C["muted"])
        etree.SubElement(rpr, qn("a:latin"), typeface=self.F["latin"])
        etree.SubElement(rpr, qn("a:ea"), typeface=self.F["ea"])
        etree.SubElement(fld, qn("a:t")).text = str(len(self.prs.slides))
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            p.remove(end)
            p.append(end)

    def tag(self, s, text, kind):
        if not text:
            return
        y = self.body_top - 0.77
        self.textbox(s, 0.6, y, 0.3, 0.5, "/", size=18, color="dark", anchor=MSO_ANCHOR.MIDDLE)
        self.pill(s, 0.9, y, text_w(text, 17) + 0.7, 0.5, text, kind or "neutral", size=17)

    def body(self):
        return MARGIN, self.body_top, W - 2 * MARGIN, H - self.body_top - 0.5

    # ------------------------------------------------------------ visuals
    def figure(self, s, x, y, w, h, fig, label):
        fig = norm_fig(fig)
        marks = fig.get("marks") or []
        if marks and self.in_focus_seq:
            # a focus sequence walks through the figure; red boxes & co. are single-slide call-outs
            WARN.append(f"[{label}] marks on {fig.get('id') or 'figure'} are not drawn on progressive-focus pages "
                        "— put them on their own slide")
            marks = []
        n0 = len(s.shapes)
        gut = mark_gutters(marks, w, h)
        x, y = x + gut["left"], y + gut["top"]
        w, h = w - gut["left"] - gut["right"], h - gut["top"] - gut["bottom"]
        rect = self.figure_body(s, x, y, w, h, fig, label)
        if marks and rect[2] < 3.5:
            WARN.append(f"[{label}] annotated {fig.get('id') or 'figure'} is only {rect[2]:.1f}in wide — "
                        "use balanced/visual density or a smaller text_ratio")
        if marks:
            self.marks(s, rect, {**fig, "marks": marks}, gut, label)
        self.focus_figure(s, rect, fig, n0, label)
        if marks and fig.get("reveal") != "click":  # animated marks must stay ungrouped (PowerPoint rule)
            grp = s.shapes.add_group_shape(list(s.shapes)[n0:])
            grp.name = f"標註/{fig.get('id') or '圖'}"

    def focus_figure(self, s, rect, fig, n0, label):
        """On a focus page: keep one region (or the whole figure) bright and veil the rest of the screenshot."""
        if not self.focus:
            return
        regions = {r.get("id"): r for r in fig.get("regions", [])}
        fid = fig.get("id")
        rx, ry, rw, rh = rect
        name = f"聚焦/{fid or '圖'}"
        if self.focus in regions:
            geo = mark_geometry({**regions[self.focus], "type": "band"})
            if geo is None:
                WARN.append(f"[{label}] region '{self.focus}' of {fid} needs x/y/w/h or row/col")
                return
            self.focus_hit = True
            fx, fy, fw, fh = geo
            # four veils around the lit region, clipped to the image
            for k, (vx, vy, vw, vh) in enumerate([(0, 0, 1, fy), (0, fy + fh, 1, 1 - fy - fh),
                                                   (0, fy, fx, fh), (fx + fw, fy, 1 - fx - fw, fh)]):
                if vw > 0.001 and vh > 0.001:
                    v = self.shape(s, MSO_SHAPE.RECTANGLE, rx + vx * rw, ry + vy * rh, vw * rw, vh * rh,
                                   fill="FFFFFF")
                    set_alpha(v, 100 - FADE_ALPHA)
                    v.name = f"{name}/遮罩{k + 1}"
        elif self.focus == fid:
            self.focus_hit = True  # the whole figure is the lit block
        else:  # another figure (or diagram part) is lit: fade this whole figure
            for sh in list(s.shapes)[n0:]:
                fade(sh)

    def figure_body(self, s, x, y, w, h, fig, label):
        """Draw the screenshot (or a placeholder) inside the cell; returns its rect (x, y, w, h)."""
        path = fig.get("path")
        cap = self.plain(f"{fig.get('id', '')}｜{fig.get('caption', '')}".strip("｜"))
        slot = self.concept_slot(fig.get("concept"), label)
        if path:
            p = Path(path)
            p = p if p.is_absolute() else self.base / p
            if p.exists() and p.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp"):
                from PIL import Image
                ch = 0.35 if cap else 0
                with Image.open(p) as im:
                    ar = im.width / im.height
                iw, ih = w, w / ar
                if ih > h - ch:
                    ih = h - ch
                    iw = ih * ar
                ix, iy = x + (w - iw) / 2, y + (h - ch - ih) / 2
                s.shapes.add_picture(str(p), Inches(ix), Inches(iy), Inches(iw), Inches(ih))
                if slot:  # concept frame around the screenshot
                    fr = self.shape(s, MSO_SHAPE.RECTANGLE, ix - 0.05, iy - 0.05, iw + 0.1, ih + 0.1,
                                    line=f"concept{slot}_fg")
                    fr.line.width = Pt(2.25)
                    fr.name = f"概念框/{fig['concept']}"
                if cap:
                    self.textbox(s, x, iy + ih + 0.08, w, 0.3, cap, size=11, color="muted",
                                 align=PP_ALIGN.CENTER)
                return ix, iy, iw, ih
            WARN.append(f"[{label}] image missing or unsupported (use PNG/JPG): {path} — placeholder drawn")
        sh = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
                        fill=f"concept{slot}_bg" if slot else "ph_bg",
                        line=f"concept{slot}_fg" if slot else "ph_line", radius=0.04, dash=True)
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, (t, sz, col, b) in enumerate([(cap or "Figure", 15, "ph_text", True),
                                             ("從論文 PDF 截圖後替換", 11, "muted", False)]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = t
            self.style(r, sz, col, b)
        return x, y, w, h

    # ------------------------------------------------------------ screenshot marks
    def marks(self, s, rect, fig, gut, label):
        """Overlay box / band / note marks on a figure. Coordinates are 0-1 fractions of the image."""
        rx, ry, rw, rh = rect
        name = f"標註/{fig.get('id') or '圖'}"
        reveal = fig.get("reveal") == "click"
        side_items = {"left": [], "right": [], "top": [], "bottom": []}  # outside labels, placed after
        groups = []  # shape ids per mark, for click steps
        for k, m in enumerate(fig["marks"], start=1):
            kind = m.get("type")
            geo = mark_geometry(m)
            if geo is None or kind not in ("box", "band", "note"):
                WARN.append(f"[{label}] {name} mark {k}: needs type box/band/note and x/y (or row/col) — skipped")
                continue
            fx, fy, fw, fh = geo
            if not (0 <= fx <= 1 and 0 <= fy <= 1 and fx + fw <= 1.001 and fy + fh <= 1.001):
                WARN.append(f"[{label}] {name} mark {k}: coordinates must be 0–1 fractions of the image")
            ax, ay, aw, ah = rx + fx * rw, ry + fy * rh, fw * rw, fh * rh
            ids = []
            if kind == "box":
                b = self.shape(s, MSO_SHAPE.RECTANGLE, ax, ay, aw, ah, line=m.get("color", "mark"))
                b.line.width = Pt(2)
                b.name = f"{name}/紅框{k}"
                ids.append(b.shape_id)
            elif kind == "band":
                slot = self.concept_slot(m.get("concept"), label)
                bg, fg = (f"concept{slot}_bg", f"concept{slot}_fg") if slot else ("inactive", "active")
                b = self.shape(s, MSO_SHAPE.RECTANGLE, ax, ay, aw, ah, fill=bg)
                clr = b.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                etree.SubElement(clr, qn("a:alpha")).set("val", str(BAND_ALPHA * 1000))
                b.name = f"{name}/色塊{k}"
                ids.append(b.shape_id)
                if m.get("text"):
                    side = m.get("label_side", "right")
                    if side == "inside":
                        t = self.textbox(s, ax + 0.08, ay, max(aw - 0.16, 0.5), ah, m["text"], size=13, color=fg,
                                         bold=True, anchor=MSO_ANCHOR.MIDDLE)
                        t.name = f"{name}/色塊{k}說明"
                        ids.append(t.shape_id)
                    elif side in side_items:
                        side_items[side].append({"k": k, "text": m["text"], "color": fg, "ids": ids,
                                                 "at": (ax + aw / 2, ay + ah / 2), "arrow": False,
                                                 "tag": f"色塊{k}說明"})
                    else:
                        WARN.append(f"[{label}] {name} mark {k}: label_side must be one of {NOTE_SIDES}")
            else:  # note: arrow from a one-line explanation to the point (x, y)
                side = m.get("side", "right")
                item = {"k": k, "text": m.get("text", ""), "color": m.get("text_color", "dark"), "ids": ids,
                        "at": (ax, ay), "arrow": m.get("color", "mark"), "tag": f"說明{k}"}
                if side == "inside":
                    tx = rx + m.get("tx", min(fx + 0.1, 0.7)) * rw
                    ty = ry + m.get("ty", max(fy - 0.15, 0.0)) * rh
                    self.note_label(s, item, tx, ty, 2.2, name)
                elif side in side_items:
                    side_items[side].append(item)
                else:
                    WARN.append(f"[{label}] {name} mark {k}: side must be one of {NOTE_SIDES}")
            groups.append((m, ids))
        self.place_side_labels(s, rect, gut, side_items, name)
        self.mark_units += sum(1 for m, _ in groups if not m.get("with_previous"))
        if reveal:
            for m, ids in groups:
                self.add_step(ids, merge=bool(m.get("with_previous")))
                if m.get("exit"):  # appear on this click, disappear on the next one (one group, not two)
                    self.pending_out = self.pending_out + ids
        elif any(m.get("exit") for m, _ in groups):
            WARN.append(f"[{label}] {name}: 'exit' needs \"reveal\": \"click\" on the figure — ignored")

    def note_label(self, s, item, tx, ty, tw, name, align=PP_ALIGN.LEFT, anchor_pt=None):
        """Text box at (tx, ty) plus, for notes, an arrow from the box edge to the target point."""
        th = 0.34 * max(1, -(-label_w(item["text"]) // max(tw, 0.5)))
        t = self.textbox(s, tx, ty, tw, th, item["text"], size=13, color=item["color"], bold=True, align=align)
        t.name = f"{name}/{item['tag']}"
        item["ids"].append(t.shape_id)
        if item["arrow"]:
            px, py = item["at"]
            sx, sy = anchor_pt or (tx if px < tx else tx + tw, ty + th / 2)
            a = self.line(s, sx, sy, px, py, item["arrow"], 1.75, arrow=True)
            a.name = f"{name}/{item['tag']}箭頭"
            item["ids"].append(a.shape_id)
        return th

    def place_side_labels(self, s, rect, gut, side_items, name):
        rx, ry, rw, rh = rect
        for side, items in side_items.items():
            if not items:
                continue
            if side in ("left", "right"):
                tw = gut[side] - 0.35
                tx = rx + rw + 0.3 if side == "right" else rx - gut[side] + 0.05
                align = PP_ALIGN.LEFT if side == "right" else PP_ALIGN.RIGHT
                floor = ry - 0.2
                for it in sorted(items, key=lambda i: i["at"][1]):
                    th = 0.34 * max(1, -(-label_w(it["text"]) // max(tw, 0.5)))
                    ty = max(it["at"][1] - th / 2, floor)
                    edge = tx if side == "right" else tx + tw
                    self.note_label(s, it, tx, ty, tw, name, align, anchor_pt=(edge, ty + th / 2))
                    floor = ty + th + 0.08
            else:
                ty = ry - gut["top"] if side == "top" else ry + rh + 0.4  # below the caption
                right_edge = rx - 0.5
                for it in sorted(items, key=lambda i: i["at"][0]):
                    tw = min(label_w(it["text"]) + 0.1, max(rw / len(items), 1.2))
                    tx = max(it["at"][0] - tw / 2, right_edge + 0.1, rx - 0.3)
                    anchor = (tx + tw / 2, ty + 0.34 if side == "top" else ty)
                    self.note_label(s, it, tx, ty, tw, name, PP_ALIGN.CENTER, anchor_pt=anchor)
                    right_edge = tx + tw

    def figures(self, s, x, y, w, h, figs, label, stack=False):
        n = len(figs)
        gap = 0.35
        if n == 0:
            return
        if stack and n <= 2:
            cells = [(x, y + k * ((h - gap * (n - 1)) / n + gap), w, (h - gap * (n - 1)) / n) for k in range(n)]
        elif n == 4 or (stack and n > 2):
            cw, ch = (w - gap) / 2, (h - gap) / 2
            cells = [(x + (k % 2) * (cw + gap), y + (k // 2) * (ch + gap), cw, ch) for k in range(min(n, 4))]
        else:
            n = min(n, 3)
            cw = (w - gap * (n - 1)) / n
            cells = [(x + k * (cw + gap), y, cw, h) for k in range(n)]
        if len(figs) > len(cells):
            WARN.append(f"[{label}] only {len(cells)} figures fit; move the rest to another slide")
        for f, cell in zip(figs, cells):
            self.figure(s, *cell, f, label)

    def table(self, s, x, y, w, rows, size=15):
        shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                 Inches(0.55 * len(rows)))
        for r, row in enumerate(rows):
            # a row whose first cell starts with {{concept}} is filled with that concept's colour
            first = parse_marks(str(row[0])) if r else []
            row_concept = first[0][2] if first and first[0][2] in self.concepts else None
            for c, val in enumerate(row):
                cell = shp.table.cell(r, c)
                cell.fill.solid()
                bg = "active" if r == 0 else ("FFFFFF" if r % 2 else "tint")
                if row_concept:
                    bg = f"concept{self.concepts[row_concept]}_bg"
                cell.fill.fore_color.rgb = self.rgb(bg)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
                for seg, hl, concept in self.segments(val):
                    rr = p.add_run()
                    rr.text = seg
                    col = "FFFFFF" if r == 0 else ("todo" if "待填" in seg else "dark")
                    self.style(rr, size, col, bold=r == 0 or hl or concept is not None)
                    if concept is not None and not row_concept and r:
                        rr.font.color.rgb = self.rgb(f"concept{self.concepts[concept]}_fg")

    def conclusion_pos(self, d, label=None):
        pos = d.get("conclusion_pos") or self.spec.get("style", {}).get("conclusion_pos", "center")
        if pos not in ("center", "bottom"):
            if label:
                WARN.append(f"[{label}] unknown conclusion_pos '{pos}' — using 'center'")
            pos = "center"
        return pos

    def conclusion(self, s, d, label):
        """Translucent white veil over the content under the heading, plus a full-width conclusion bar.

        pos 'center': the veil covers the whole body and the bar sits in its middle (no space reserved).
        pos 'bottom': space is reserved under the content; the veil stops at the bar at the bottom.
        Both fade in together on one click, unless building keyframes (then this slide is the 'after' frame).
        """
        text, kind = self.plain(d["conclusion"]), d.get("conclusion_kind", "problem")
        if kind not in ("problem", "solution", "neutral"):
            WARN.append(f"[{label}] unknown conclusion_kind '{kind}' — using 'problem'")
            kind = "problem"
        if kind in ("problem", "solution"):
            self.use(kind, text)
        x, w, top = 0.6, W - 1.2, self.body_top - 0.15
        if self.conclusion_pos(d, label) == "bottom":
            y = H - 0.5 - BAR_H
            veil_h = y - top
        else:
            veil_h = H - 0.5 - top
            y = top + (veil_h - BAR_H) / 2
        veil = self.shape(s, MSO_SHAPE.RECTANGLE, x, top, w, veil_h, fill="FFFFFF")
        veil.name = "結論遮罩"
        clr = veil.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha")).set("val", str(VEIL_ALPHA * 1000))

        size = 20
        while size > 14 and text_w(text, size) > w - 0.6:
            size -= 1
        if text_w(text, size) > w - 0.6:
            WARN.append(f"[{label}] conclusion bar text may overflow at {size}pt — shorten it")
        bar = self.shape(s, MSO_SHAPE.RECTANGLE, x, y, w, BAR_H, fill=kind)
        bar.name = "結論橫條"
        tf = bar.text_frame
        tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        self.style(r, size, "FFFFFF", bold=True)
        self.add_step([veil.shape_id, bar.shape_id])

    def key_point(self, s, x, y, w, h, text):
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill="tint", radius=0.12)
        self.textbox(s, x + 0.3, y, w - 0.6, h, [{"t": "本頁重點", "size": 12, "color": "active", "bold": True},
                                                 {"t": text, "size": 17, "color": "dark", "bold": True}],
                     anchor=MSO_ANCHOR.MIDDLE, space=2)

    def notes(self, s, paras):
        if paras:
            s.notes_slide.notes_text_frame.text = "\n".join(p if isinstance(p, str) else str(p) for p in paras)

    # ------------------------------------------------------------ slide kinds
    def s_cover(self, d):
        style = d.get("style") or ("paper" if self.nav_style == "double" else "centered")
        if style not in ("centered", "paper"):
            WARN.append(f"[cover] unknown cover.style '{style}' — using 'centered'")
            style = "centered"
        s = self.slide()
        (self.cover_paper if style == "paper" else self.cover_centered)(s, d)
        self.notes(s, d.get("notes"))

    def cover_centered(self, s, d):
        """Single-nav cover: centred title, orange rule, then 作者．用途."""
        self.textbox(s, 1.2, 1.9, W - 2.4, 1.5, d.get("title", ""), size=38, color="dark", bold=True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, min_size=28, label="cover")
        self.line(s, 3.2, 3.6, W - 3.2, 3.6)
        y = 3.85
        for key, size, color in (("subtitle", 15, "muted"), ("byline", 17, "text"), ("date", 14, "muted")):
            if d.get(key):
                self.textbox(s, 1.2, y, W - 2.4, 0.5, d[key], size=size, color=color, align=PP_ALIGN.CENTER)
                y += 0.55

    def cover_paper(self, s, d):
        """Double-nav cover: venue badge, serif paper title, orange rule, authors | affiliations, presenter."""
        x, w = MARGIN, W - 2 * MARGIN
        if d.get("venue"):
            self.pill(s, x, 1.35, text_w(d["venue"], 14) + 0.5, 0.42, d["venue"], "active", size=14)
        self.textbox(s, x, 1.95, w, 1.65, d.get("paper_title") or d.get("title", ""), size=34, color="dark",
                     bold=True, anchor=MSO_ANCHOR.BOTTOM, latin=self.F["ref_latin"], min_size=24,
                     label="cover")
        self.line(s, x, 3.8, W - MARGIN, 3.8)
        # no author list (e.g. a progress report) -> fall back to the byline so the cover is not empty
        cols = [(as_list(d.get("authors") or d.get("byline")), "dark"), (as_list(d.get("affiliations")), "muted")]
        cw = (w - 0.5) / 2
        for k, (items, color) in enumerate(cols):
            if items:
                self.textbox(s, x + k * (cw + 0.5), 4.05, cw, 1.8, items, size=16, color=color, space=4,
                             latin=self.F["ref_latin"], min_size=12, label="cover")
        if d.get("presenter"):
            self.textbox(s, W - MARGIN - 5, H - 1.05, 5, 0.45, f"報告者：{d['presenter']}", size=17,
                         color="text", align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        if d.get("date"):
            self.textbox(s, x, H - 1.05, 4, 0.45, d["date"], size=14, color="muted", anchor=MSO_ANCHOR.MIDDLE)

    def s_content(self, d, label):
        dens = DENSITIES[d.get("density", self.density)]
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d)
        x, y, w, h = self.body()
        if d.get("conclusion") and self.conclusion_pos(d) == "bottom":
            h -= BAR_H + 0.15  # reserved on the 'before' keyframe too, so nothing jumps between frames
        key = dens["key"]
        txt = pick_text(d, key)
        figs, table = d.get("figs", []), d.get("table")
        has_vis = bool(figs or table)
        is_list = txt is not None and txt is d.get("short") and isinstance(txt, list)

        if key == "one_line" or not txt:
            line_txt = pick_text(d, "one_line")
            if isinstance(line_txt, list):
                line_txt = line_txt[0]
            if not has_vis:  # statement slide
                self.textbox(s, x + 0.8, y, w - 1.6, h - 0.3, line_txt or "", size=30, color="dark", bold=True,
                             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=22, label=label)
            else:
                if line_txt:
                    self.textbox(s, x, y - 0.05, w, 0.5, line_txt, size=dens["size"], color="dark", bold=True,
                                 min_size=14, label=label)
                vy, vh = y + 0.65, h - 0.65
                if table and not figs:
                    self.table(s, x + 1.2, vy + 0.2, w - 2.4, table)
                else:
                    self.figures(s, x, vy, w, vh, figs, label)
        else:
            tw = w * (d.get("text_ratio") or dens["text_ratio"]) if has_vis else w
            kp = dens["key_point"] and d.get("one_line") and key == "full"
            th = h - 1.3 if kp else h
            self.textbox(s, x, y + 0.1, tw, th - 0.1, txt, size=dens["size"], bullet=is_list,
                         space=14 if is_list else 12, min_size=dens["min"], label=label)
            if kp:
                self.key_point(s, x, y + h - 1.0, tw, 1.0, d["one_line"])
            if has_vis:
                vx, vw = x + tw + 0.5, w - tw - 0.5
                if table:
                    self.table(s, vx, y + 0.2, vw, table, size=14 if vw < 6 else 15)
                    # figs beside a table are the visual-density alternative, not drawn here
                else:
                    self.figures(s, vx, y, vw, h, figs, label, stack=True)
        if d.get("callout"):
            self.textbox(s, x, H - 0.45, w, 0.3, d["callout"], size=12, color="muted")
        if d.get("conclusion"):
            self.conclusion(s, d, label)
        notes = d.get("notes") or (as_list(d.get("full")) if dens["auto_notes"] else [])
        self.notes(s, notes)

    def s_mapping(self, d, label):
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d, "neutral")
        pairs = d["pairs"]
        pw, ph = 3.6, 0.62
        lx, rx = 2.1, W - 2.1 - pw
        top = self.body_top + 0.1
        step = min(1.45, (H - top - 0.9) / max(len(pairs), 1))
        for k, (a, b) in enumerate(pairs):
            yy = top + k * step
            size = 17 if max(text_w(a, 17), text_w(b, 17)) < pw - 0.4 else 14
            self.pill(s, lx, yy, pw, ph, a, "problem", size=size, source="mapping")
            self.pill(s, rx, yy, pw, ph, b, "solution", size=size, source="mapping")
            self.line(s, lx + pw + 0.25, yy + ph / 2, rx - 0.25, yy + ph / 2, "444444", 1.75, arrow=True)
        self.notes(s, d.get("notes"))

    def s_split(self, d, label):
        dens = DENSITIES[d.get("density", self.density)]
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d, "neutral")
        x, y, w, _ = self.body()
        brief = dens["key"] == "one_line"
        h = 3.4 if brief else 3.9
        pw = (w - 0.8) / 2
        sides = [(d.get("good_title", "優點"), d.get("good", []), "good_bg", "good_fg", PP_ALIGN.LEFT),
                 (d.get("bad_title", "限制"), d.get("bad", []), "bad_bg", "bad_fg", PP_ALIGN.RIGHT)]
        for k, (title, items, bg, fg, align) in enumerate(sides):
            px = x + k * (pw + 0.8)
            self.shape(s, MSO_SHAPE.RECTANGLE, px, y, pw, h, fill=bg)
            self.textbox(s, px + 0.35, y + 0.3, pw - 0.7, 0.45, title, size=15, color="muted", bold=True,
                         align=align)
            paras = []
            for it in items:
                t, desc = (it + [""])[:2] if isinstance(it, list) else (it, "")
                paras.append({"t": t, "size": 19 if brief else 17, "color": fg, "bold": True})
                if desc and not brief:
                    paras.append({"t": desc, "size": 15, "color": "text"})
                    paras.append({"t": " ", "size": 8})
            self.textbox(s, px + 0.35, y + 0.95, pw - 0.7, h - 1.2, paras, align=align,
                         space=16 if brief else 6, min_size=12, size=17, label=label)
        self.line(s, W / 2, y + 0.2, W / 2, y + h - 0.2)
        auto = [f"{(it + [''])[0]}：{(it + [''])[1]}" for it in d.get("good", []) + d.get("bad", [])
                if isinstance(it, list)] if brief else []
        self.notes(s, d.get("notes") or auto)

    def s_refs(self, d, label):
        brief = DENSITIES[d.get("density", self.density)]["key"] == "one_line"
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d, "neutral")
        paras = []
        for r in d["refs"]:
            r = r if isinstance(r, dict) else {"title": r[0], "source": r[1] if len(r) > 1 else "",
                                                "gist": r[2] if len(r) > 2 else ""}
            paras.append({"t": f"–  {r['title']}", "size": 17, "color": "dark", "bold": True, "space": 3})
            if r.get("source"):
                paras.append({"t": f"    {r['source']}", "size": 13, "color": "muted", "space": 3})
            if r.get("gist") and not brief:
                paras.append({"t": f"    {r['gist']}", "size": 14, "color": "text", "space": 3})
            paras.append({"t": " ", "size": 8, "space": 3})
        top = self.body_top if d.get("title" if self.nav_style == "double" else "tag") else self.body_top - 0.45
        self.textbox(s, 1.2, top, W - 2.4, H - top - 0.4, paras, latin=self.F["ref_latin"], size=17,
                     min_size=11, label=label)
        self.notes(s, d.get("notes"))

    def s_statement(self, d, label):
        s = self.slide()
        if d.get("chapter"):
            self.nav(s, d.get("chapter"), d.get("section"))
        self.textbox(s, 1.5, 1.8, W - 3, 4.2, d.get("text", ""), size=40, color="dark", bold=True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=24, label=label)
        self.notes(s, d.get("notes"))

    # ------------------------------------------------------------ diagram (the presenter's own concept diagram)
    def s_diagram(self, d, label):
        """Full-page native diagram: nodes on a col/row grid, elbow connectors glued to nodes, module frames."""
        spec = d.get("diagram")
        if isinstance(spec, str):
            name, spec = spec, (self.spec.get("diagrams") or {}).get(spec)
            if spec is None:
                WARN.append(f"[{label}] diagram '{name}' is not defined in top-level diagrams — skipped")
                return
        spec = spec or {}
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d)
        top = self.body_top
        bottom = H - 0.6
        if d.get("conclusion") and self.conclusion_pos(d) == "bottom":
            bottom = H - 0.5 - BAR_H - 0.2
        DiagramLayout(self, s, spec, d, label).draw(MARGIN, top, W - 2 * MARGIN, bottom - top)
        if d.get("conclusion"):
            self.conclusion(s, d, label)
        self.notes(s, d.get("notes") or as_list(d.get("full") or d.get("short")))

    # ------------------------------------------------------------ timeline
    def s_timeline(self, d, label):
        """Years on a bottom axis, one horizontal band per category, items as small pills (years only)."""
        s = self.slide()
        self.nav(s, d.get("chapter"), d.get("section"))
        self.heading(s, d)
        items = [dict(it) for it in d.get("items", []) if self.timeline_item_ok(it, label)]
        lanes = d.get("lanes") or [{"name": ""}]
        names = [ln.get("name", "") for ln in lanes]
        for it in items:
            it.setdefault("lane", names[0])
            if it["lane"] not in names:
                WARN.append(f"[{label}] timeline item '{it.get('label')}' uses unknown lane '{it['lane']}'")
        items = [it for it in items if it["lane"] in names]
        if not items:
            WARN.append(f"[{label}] timeline has no items")
            return
        y0, y1 = d.get("years") or (min(it["year"] for it in items), max(it["year"] for it in items))
        ncol = y1 - y0 + 1

        bottom = H - 0.6 - 0.55  # leave room for the year axis
        if d.get("conclusion") and self.conclusion_pos(d) == "bottom":
            bottom -= BAR_H + 0.2
        top = self.body_top
        lab_w = max([label_w(self.plain(n)) + 0.45 for n in names if n] + [0]) if any(names) else 0
        lab_w = min(max(lab_w, 1.3), 2.4) if lab_w else 0
        x0, x1 = MARGIN, W - MARGIN
        gx0, gx1 = x0 + lab_w + 0.15, x1 - 0.15  # the year grid
        colw = (gx1 - gx0) / ncol
        band_gap = 0.12
        xs = lambda yr: gx0 + (yr - y0 + 0.5) * colw
        plans = {ln.get("name", ""): self.timeline_plan([it for it in items if it["lane"] == ln.get("name", "")], xs)
                 for ln in lanes}
        band_h = (bottom - top - band_gap * (len(lanes) - 1)) / len(lanes)
        # bands only as tall as their content needs (so a sparse timeline doesn't float in empty colour)
        need = max(pl["height"] for pl in plans.values()) + 0.7
        band_h = min(band_h, max(1.3, need))
        bottom = top + len(lanes) * band_h + band_gap * (len(lanes) - 1)
        drawn = []
        for k, ln in enumerate(lanes):
            by = top + k * (band_h + band_gap)
            slot = self.concept_slot(ln.get("concept"), label)
            bg, fg = (f"concept{slot}_bg", f"concept{slot}_fg") if slot else ("lane_bg", "lane_fg")
            band = self.shape(s, MSO_SHAPE.RECTANGLE, x0, by, x1 - x0, band_h, fill=bg)
            band.name = f"時間軸/分類/{self.plain(ln.get('name', '')) or k + 1}"
            drawn.append(band)
            if ln.get("name"):
                t = self.textbox(s, x0 + 0.2, by, lab_w - 0.2, band_h, ln["name"], size=15, color=fg, bold=True,
                                 anchor=MSO_ANCHOR.MIDDLE)
                t.name = f"{band.name}/名稱"
                drawn.append(t)
            drawn += self.timeline_lane(s, plans[ln.get("name", "")], xs, by, band_h, fg, label)
        # year axis
        ay = bottom + 0.2
        ax = self.line(s, gx0 - 0.1, ay, gx1, ay, "muted", 1.25)
        ax.name = "時間軸/年份軸"
        drawn.append(ax)
        every = max(1, int(-(-0.55 // colw)))  # skip labels when columns get narrow
        for yr in range(y0, y1 + 1):
            tick = self.line(s, xs(yr), ay - 0.06, xs(yr), ay + 0.06, "muted", 1.25)
            tick.name = f"時間軸/年份軸/{yr}"
            drawn.append(tick)
            if (yr - y0) % every == 0:
                t = self.textbox(s, xs(yr) - 0.4, ay + 0.1, 0.8, 0.3, str(yr), size=12, color="muted",
                                 align=PP_ALIGN.CENTER)
                t.name = f"時間軸/年份軸/{yr}/標籤"
                drawn.append(t)
        grp = s.shapes.add_group_shape(drawn)
        grp.name = "時間軸"
        if d.get("conclusion"):
            self.conclusion(s, d, label)
        self.notes(s, d.get("notes") or as_list(d.get("full") or d.get("short")))

    def timeline_item_ok(self, it, label):
        if not isinstance(it.get("year"), int):
            WARN.append(f"[{label}] timeline item '{it.get('label')}' needs an integer year (years only) — skipped")
            return False
        return True

    TL_PILL_H, TL_NOTE_H, TL_GAP = 0.38, 0.24, 0.1

    def timeline_plan(self, items, xs):
        """Assign items to tracks so neighbours in the same band never overlap horizontally."""
        tracks = []  # each: right edge of the last pill
        placed = []
        for it in sorted(items, key=lambda i: (i["year"], i.get("label", ""))):
            w = label_w(self.plain(it.get("label", ""))) + 0.4
            left = xs(it["year"]) - w / 2
            if it.get("note"):
                w_note = text_w(self.plain(it["note"]), 11) + 0.1
                left = min(left, xs(it["year"]) - w_note / 2)
                w = max(w, w_note)
            k = next((i for i, r in enumerate(tracks) if left >= r + 0.08), None)
            if k is None:
                tracks.append(0)
                k = len(tracks) - 1
            tracks[k] = left + w
            placed.append((it, k))
        row_h = self.TL_PILL_H + (self.TL_NOTE_H if any(it.get("note") for it in items) else 0) + self.TL_GAP
        return {"placed": placed, "tracks": len(tracks), "row_h": row_h,
                "height": max(len(tracks) * row_h - self.TL_GAP, 0), "n": len(items)}

    def timeline_lane(self, s, plan, xs, by, band_h, fg, label):
        """Draw a lane's pills; its tracks are centred vertically in the band."""
        pill_h, note_h = self.TL_PILL_H, self.TL_NOTE_H
        placed, row_h, total = plan["placed"], plan["row_h"], plan["height"]
        if total > band_h - 0.1:
            WARN.append(f"[{label}] timeline band is too crowded ({plan['n']} items in {plan['tracks']} rows) — "
                        "split the timeline or widen the year range")
        start = by + (band_h - total) / 2
        out = []
        for it, k in placed:
            new = bool(it.get("highlight"))
            text = self.plain(it.get("label", ""))
            w = label_w(text) + 0.4
            y = start + k * row_h
            pill = self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, xs(it["year"]) - w / 2, y, w, pill_h,
                              fill="new_bg" if new else "FFFFFF", line="new_line" if new else fg, radius=0.5)
            pill.line.width = Pt(2.25 if new else 1.25)
            tf = pill.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            r = para.add_run()
            r.text = text
            self.style(r, 13, "dark", True)
            pill.name = f"時間軸/{text}"
            if new:
                self.use("new", text)
            out.append(pill)
            if it.get("note"):
                t = self.textbox(s, xs(it["year"]) - 0.9, y + pill_h + 0.02, 1.8, note_h, it["note"], size=11,
                                 color="muted", align=PP_ALIGN.CENTER)
                t.name = f"時間軸/{text}/註記"
                out.append(t)
        return out

    def build(self):
        self.s_cover(self.spec.get("cover", {}))
        for i, d in enumerate(self.spec.get("slides", []), start=2):
            t = d.get("type", "content")
            label = f"slide {i} {d.get('tag') or t}"
            fn = getattr(self, f"s_{t}", None) if t != "cover" else None
            if fn is None:
                hint = " (divider pages were removed — delete this slide)" if t == "divider" else ""
                WARN.append(f"[{label}] unknown slide type '{t}' — skipped{hint}")
                continue
            for page in self.focus_pages(d, label):
                self.render_page(fn, page, label)
        WARN[:] = list(dict.fromkeys(WARN))  # rebuilt keyframes repeat their warnings
        return self.prs

    def focus_pages(self, d, label):
        """A slide with `focus` expands into an overview page (unless focus_overview is false) plus one page per
        step; each step lights one id and may override the slide's text. `conclusion` is not inherited by steps."""
        if not d.get("focus"):
            return [d]
        base = {k: v for k, v in d.items() if k not in ("focus", "focus_overview")}
        pages = [base] if d.get("focus_overview", True) else []
        inherit = {k: v for k, v in base.items() if k not in ("conclusion", "conclusion_kind")}
        for k, st in enumerate(d["focus"], start=1):
            on = st.get("on")
            if isinstance(on, list):
                WARN.append(f"[{label}] focus step {k}: one block per step — using '{on[0]}'")
                on = on[0] if on else None
            if not on:
                WARN.append(f"[{label}] focus step {k} has no 'on' — skipped")
                continue
            page = {**inherit, **{x: st[x] for x in FOCUS_KEYS if x in st}, "_focus": on, "_focus_seq": True}
            pages.append(page)
        if pages and pages[0] is base:
            pages[0] = {**base, "_focus_seq": True}
        return pages

    def render_page(self, fn, d, label):
        self.focus, self.focus_hit = d.get("_focus"), False
        self.in_focus_seq = bool(d.get("_focus_seq"))
        fn(d, label)
        self.close_steps()
        if self.focus and not self.focus_hit:
            WARN.append(f"[{label}] focus '{self.focus}' matches no diagram node/module or figure region here")
        steps = self.steps
        if self.mark_units > MAX_MARKS:
            WARN.append(f"[{label}] {self.mark_units} annotations (red boxes / bands / notes) on one slide — "
                        f"keep it to {MAX_MARKS}; split the slide")
        if steps and self.keyframes:
            # one slide per click: frame k shows what is visible after k clicks (the slide is rebuilt identically)
            def hidden_after(k):
                later_in = [i for st in steps[k:] for i in st["in"]]
                gone = [i for st in steps[:k] for i in st["out"]]
                return later_in + gone
            drop_shapes(self.prs.slides[-1], hidden_after(0))
            self.page_number(self.prs.slides[-1])
            for k in range(1, len(steps) + 1):
                fn(d, label)
                self.close_steps()
                drop_shapes(self.prs.slides[-1], hidden_after(k))
                self.page_number(self.prs.slides[-1])
        else:
            if steps:
                click_steps(self.prs.slides[-1], steps)
            self.page_number(self.prs.slides[-1])
        self.focus = None
        self.in_focus_seq = False

    def close_steps(self):
        """A mark that should disappear after the last click gets one more click of its own."""
        if self.pending_out:
            self.steps.append({"in": [], "out": self.pending_out})
            self.pending_out = []


    def color_table(self):
        """Rows for the report's colour/concept table: only colours this deck actually uses."""
        rows = []
        for key, label, hexes, what in (("problem", "紅", [self.C["problem"]], "問題"),
                                        ("solution", "綠", [self.C["solution"]], "解法")):
            u = self.uses.get(key)
            if u:
                names = u["mapping"] or u["names"]
                rows.append({"color": label, "hex": hexes, "meaning": f"{what}：" + "／".join(names),
                             "pages": sorted(u["pages"])})
        u = self.uses.get("new")
        if u:
            rows.append({"color": "綠色描邊", "hex": [self.C["new_line"], self.C["new_bg"]],
                         "meaning": "新增、有變化或要強調的元件（架構圖、時間軸）：" + "／".join(u["names"]),
                         "pages": sorted(u["pages"])})
        for name, slot in sorted(self.concepts.items(), key=lambda kv: kv[1]):
            u = self.uses.get(f"concept:{name}")
            if u:
                rows.append({"color": f"概念色 {slot}（{CONCEPT_NAMES[slot]}）",
                             "hex": [self.C[f"concept{slot}_bg"], self.C[f"concept{slot}_fg"]],
                             "meaning": name, "pages": sorted(u["pages"])})
        return rows


# ---------------------------------------------------------------- diagram layout
def span(v):
    """1 -> (1, 1); [2, 3] -> (2, 3)"""
    if isinstance(v, (list, tuple)):
        return int(v[0]), int(v[-1])
    return int(v), int(v)


class DiagramLayout:
    """Grid layout for a diagram spec: {direction, nodes, edges, modules}. Coordinates in inches."""

    SITE = {"top": 0, "left": 1, "bottom": 2, "right": 3}  # connection sites of rect / roundRect

    def __init__(self, deck, slide, spec, d, label):
        self.deck, self.s, self.label = deck, slide, label
        self.direction = spec.get("direction", "right")
        hide = set(d.get("hide", []))
        relabel = d.get("relabel", {})
        self.highlight = set(d.get("highlight", []))
        self.all_nodes = spec.get("nodes", [])  # the grid is sized from every node, hidden or not
        self.nodes = [{**n, "label": relabel.get(n["id"], n.get("label", n["id"]))}
                      for n in self.all_nodes if n.get("id") not in hide]
        ids = {n["id"] for n in spec.get("nodes", [])}
        for x in self.highlight | hide | set(relabel):
            if x not in ids and x not in {m.get("id") for m in spec.get("modules", [])}:
                WARN.append(f"[{label}] diagram has no node or module '{x}'")
        shown = {n["id"] for n in self.nodes}
        self.edges = []
        for e in spec.get("edges", []):
            e = e if isinstance(e, dict) else {"from": e[0], "to": e[1], "label": e[2] if len(e) > 2 else None}
            if e["from"] not in ids or e["to"] not in ids:
                WARN.append(f"[{label}] diagram edge {e['from']}→{e['to']} refers to an unknown node — skipped")
            elif e["from"] in shown and e["to"] in shown:
                self.edges.append(e)
        self.modules_all = spec.get("modules", [])
        self.modules = [m for m in self.modules_all if any(i in shown for i in m.get("nodes", []))]
        self.box = {}  # id -> (x, y, w, h)
        self.shape = {}  # id -> pptx shape

    def draw(self, ax, ay, aw, ah):
        if not self.nodes:
            WARN.append(f"[{self.label}] diagram has no nodes")
            return
        ncols = max(span(n.get("col", 1))[1] for n in self.all_nodes)
        nrows = max(span(n.get("row", 1))[1] for n in self.all_nodes)
        has_img = any(n.get("image") for n in self.all_nodes)
        framed = bool(self.modules_all)
        pad = 0.35 if framed else 0.0  # room for module frames and their titles
        ax, ay, aw, ah = ax + pad, ay + pad, aw - 2 * pad, ah - 2 * pad
        cw, ch = aw / ncols, ah / nrows
        # with frames, widen the gaps along the flow so connector bends fall between a frame and the next node
        extra = 0.45 if framed else 0
        gx = min(1.1, max(0.55, cw * 0.32)) + (extra if self.direction != "down" else 0)
        gy = min(0.9, max(0.4, ch * 0.3)) + (extra if self.direction == "down" else 0)
        nw = min(cw - gx, 3.0)
        nh = min(ch - gy, 2.4 if has_img else 1.1)
        if nw < 1.0 or nh < 0.4:
            WARN.append(f"[{self.label}] diagram grid {ncols}×{nrows} is too dense for one slide — split it")
            nw, nh = max(nw, 0.8), max(nh, 0.4)
        for n in self.all_nodes:
            c0, c1 = span(n.get("col", 1))
            r0, r1 = span(n.get("row", 1))
            w = nw + (c1 - c0) * cw
            h = nh + (r1 - r0) * ch
            cx = ax + (c0 - 1 + (c1 - c0 + 1) / 2) * cw
            cy = ay + (r0 - 1 + (r1 - r0 + 1) / 2) * ch
            self.box[n["id"]] = (cx - w / 2, cy - h / 2, w, h)
        drawn = []
        for m in self.modules:
            drawn += self.module(m)
        for n in self.nodes:
            drawn += self.node(n)
        for e in self.edges:
            drawn += self.edge(e)
        self.apply_focus(drawn)
        grp = self.s.shapes.add_group_shape(drawn)
        grp.name = "架構圖"

    def apply_focus(self, drawn):
        focus = self.deck.focus
        if not focus:
            return
        mod = next((m for m in self.modules if m.get("id") == focus), None)
        lit_nodes = set(mod.get("nodes", [])) if mod else ({focus} if focus in self.box else set())
        if not lit_nodes:
            return  # not ours (e.g. a figure region on another slide type)
        self.deck.focus_hit = True
        keep = {f"架構圖/{i}" for i in lit_nodes}
        keep |= {f"架構圖/{i}/縮圖" for i in lit_nodes} | {f"架構圖/{i}/文字" for i in lit_nodes}
        if mod:
            keep |= {f"架構圖/模組/{focus}", f"架構圖/模組/{focus}/標題"}
        for e in self.edges:
            if e["from"] in lit_nodes or e["to"] in lit_nodes:
                keep |= {f"架構圖/連線/{e['from']}-{e['to']}", f"架構圖/連線/{e['from']}-{e['to']}/標籤"}
        for sh in drawn:
            if sh.name not in keep:
                fade(sh)

    # -------------------------------------------------------------- parts
    def module(self, m):
        dk = self.deck
        boxes = [self.box[i] for i in m.get("nodes", []) if i in self.box]
        x0 = min(b[0] for b in boxes) - 0.2
        y0 = min(b[1] for b in boxes) - (0.42 if m.get("label") else 0.2)
        x1 = max(b[0] + b[2] for b in boxes) + 0.2
        y1 = max(b[1] + b[3] for b in boxes) + 0.2
        new = m.get("id") in self.highlight
        fr = dk.shape(self.s, MSO_SHAPE.ROUNDED_RECTANGLE, x0, y0, x1 - x0, y1 - y0,
                      fill="new_bg" if new else None, line="new_line" if new else "node_line", radius=0.06,
                      dash=True)
        fr.name = f"架構圖/模組/{m.get('id') or m.get('label', '')}"
        out = [fr]
        if new:
            dk.use("new", dk.plain(m.get("label") or m.get("id")))
        if m.get("label"):
            t = dk.textbox(self.s, x0 + 0.15, y0 + 0.06, x1 - x0 - 0.3, 0.3, m["label"], size=12,
                           color="new_line" if new else "muted", bold=True)
            t.name = f"架構圖/模組/{m.get('id') or m['label']}/標題"
            out.append(t)
        return out

    def node(self, n):
        dk, (x, y, w, h) = self.deck, self.box[n["id"]]
        slot = dk.concept_slot(n.get("concept"), self.label)
        new = n["id"] in self.highlight
        fill = "new_bg" if new else (f"concept{slot}_bg" if slot else "FFFFFF")
        line = "new_line" if new else (f"concept{slot}_fg" if slot else "node_line")
        sh = dk.shape(self.s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, radius=0.12)
        sh.line.width = Pt(2.25 if new else 1.25)
        sh.name = f"架構圖/{n['id']}"
        if new:
            dk.use("new", dk.plain(n["label"]))
        self.shape[n["id"]] = sh
        out = [sh]
        img = n.get("image")
        text_top, text_h = y, h
        if img:
            pth = Path(img)
            pth = pth if pth.is_absolute() else dk.base / pth
            ih = h * 0.62
            if pth.exists():
                from PIL import Image
                with Image.open(pth) as im:
                    ar = im.width / im.height
                pw, ph = min(w - 0.2, (ih - 0.15) * ar), min(ih - 0.15, (w - 0.2) / ar)
                pic = self.s.shapes.add_picture(str(pth), Inches(x + (w - pw) / 2), Inches(y + 0.1),
                                                Inches(pw), Inches(ph))
                pic.name = f"架構圖/{n['id']}/縮圖"
                out.append(pic)
            else:
                WARN.append(f"[{self.label}] diagram node '{n['id']}' image missing: {img} — placeholder drawn")
                ph = dk.shape(self.s, MSO_SHAPE.RECTANGLE, x + 0.15, y + 0.1, w - 0.3, ih - 0.15, fill="ph_bg",
                              line="ph_line", dash=True)
                ph.name = f"架構圖/{n['id']}/縮圖"
                out.append(ph)
            text_top, text_h = y + ih, h - ih
        size = 16
        text = dk.plain(n["label"])
        while size > 11 and est_height([{"t": text}], size, w - 0.2, 0, False) > text_h - 0.05:
            size -= 1
        if est_height([{"t": text}], size, w - 0.2, 0, False) > text_h + 0.05:
            WARN.append(f"[{self.label}] diagram node '{n['id']}' label is too long for its box")
        if img:
            t = dk.textbox(self.s, x + 0.1, text_top, w - 0.2, text_h, n["label"], size=size, color="dark",
                           bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0)
            t.name = f"架構圖/{n['id']}/文字"
            out.append(t)
        else:  # label lives in the node itself so it moves and resizes with it
            tf = sh.text_frame
            tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.08)
            tf.margin_top = tf.margin_bottom = 0
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            for seg, hl, concept in dk.segments(n["label"]):
                r = para.add_run()
                r.text = seg
                dk.style(r, size, "dark", True)
                if concept is not None:
                    r.font.color.rgb = dk.rgb(f"concept{dk.concepts[concept]}_fg")
        return out

    def edge(self, e):
        dk = self.deck
        a, b = self.box[e["from"]], self.box[e["to"]]
        (ac0, ac1), (bc0, bc1) = span(self.node_of(e["from"]).get("col", 1)), span(self.node_of(e["to"]).get("col", 1))
        (ar0, ar1), (br0, br1) = span(self.node_of(e["from"]).get("row", 1)), span(self.node_of(e["to"]).get("row", 1))
        if self.direction == "down":
            horizontal = not (br0 > ar1 or br1 < ar0)  # same row band -> sideways
        else:
            horizontal = bc0 > ac1 or bc1 < ac0  # different columns -> left/right
        if horizontal:
            fwd = b[0] > a[0]
            s1, s2 = ("right", "left") if fwd else ("left", "right")
            p1 = (a[0] + a[2] if fwd else a[0], a[1] + a[3] / 2)
            p2 = (b[0] if fwd else b[0] + b[2], b[1] + b[3] / 2)
        else:
            fwd = b[1] > a[1]
            s1, s2 = ("bottom", "top") if fwd else ("top", "bottom")
            p1 = (a[0] + a[2] / 2, a[1] + a[3] if fwd else a[1])
            p2 = (b[0] + b[2] / 2, b[1] if fwd else b[1] + b[3])
        if abs(p1[0] - p2[0]) < 0.02 or abs(p1[1] - p2[1]) < 0.02:  # aligned -> straight
            p2 = (p1[0], p2[1]) if not horizontal else (p2[0], p1[1])
            c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]),
                                            Inches(p2[0]), Inches(p2[1]))
        elif horizontal:  # horizontal-first Z: python-pptx handles the flips
            c = self.s.shapes.add_connector(MSO_CONNECTOR.ELBOW, Inches(p1[0]), Inches(p1[1]),
                                            Inches(p2[0]), Inches(p2[1]))
            set_bend(c, self.bend(e, p1[0], p2[0], s1, s2))
        else:  # vertical-first Z: a bentConnector3 rotated 90° clockwise
            c = self.s.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, Inches(1), Inches(1))
            vertical_elbow(c, p1, p2)
            set_bend(c, self.bend(e, p1[1], p2[1], s1, s2))
        c.line.color.rgb = dk.rgb(e.get("color", "edge"))
        c.line.width = Pt(1.5)
        if e.get("dashed"):
            c.line.dash_style = MSO_LINE.DASH
        etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
        etree.SubElement(c._element.spPr, qn("a:effectLst"))
        glue(c, self.shape[e["from"]], self.SITE[s1], self.shape[e["to"]], self.SITE[s2])
        c.name = f"架構圖/連線/{e['from']}-{e['to']}"
        out = [c]
        if e.get("label"):
            # on the last segment (the one entering the target), which is never shared with sibling edges
            tw = label_w(dk.plain(e["label"])) + 0.2
            if horizontal:
                mx = p2[0] - (p2[0] - (p1[0] + p2[0]) / 2) / 2
                t = dk.textbox(self.s, mx - tw / 2, p2[1] - 0.34, tw, 0.3, e["label"], size=12, color="muted",
                               align=PP_ALIGN.CENTER)
            else:
                my = p2[1] - (p2[1] - (p1[1] + p2[1]) / 2) / 2
                t = dk.textbox(self.s, p2[0] + 0.08, my - 0.15, tw, 0.3, e["label"], size=12, color="muted")
            t.name = f"架構圖/連線/{e['from']}-{e['to']}/標籤"
            out.append(t)
        return out

    def frame_margin(self, node_id, side):
        """How far a module frame reaches out from this node on the given side (0 if the node is unframed)."""
        for m in self.modules:
            if node_id in m.get("nodes", []):
                return 0.42 if side == "top" and m.get("label") else 0.2
        return 0.0

    def bend(self, e, a, b, s1, s2):
        """Fraction along the flow where a Z connector turns: the middle of the gap left free by frames."""
        a2 = a + (self.frame_margin(e["from"], s1) + 0.05) * (1 if b > a else -1)
        b2 = b - (self.frame_margin(e["to"], s2) + 0.05) * (1 if b > a else -1)
        if (b2 - a2) * (b - a) <= 0:  # frames eat the whole gap: fall back to the middle
            return 0.5
        return ((a2 + b2) / 2 - a) / (b - a)

    def node_of(self, i):
        return next(n for n in self.nodes if n["id"] == i)


def set_bend(c, frac):
    """Position of the middle segment of a bentConnector3, as a fraction of the way from its start."""
    av = c._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(c._element.spPr.find(qn("a:prstGeom")), qn("a:avLst"))
    for g in list(av):
        av.remove(g)
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj1")
    gd.set("fmla", f"val {int(round(frac * 100000))}")


def vertical_elbow(c, p1, p2):
    """Make connector c a vertical-first elbow from p1 to p2 (inches) by rotating a bentConnector3 90° cw."""
    (x1, y1), (x2, y2) = p1, p2
    dx, dy = x2 - x1, y2 - y1
    w, h = abs(dy), abs(dx)
    cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
    xfrm = c._element.spPr.find(qn("a:xfrm"))
    for k in ("flipH", "flipV"):
        xfrm.attrib.pop(k, None)
    xfrm.set("rot", "5400000")
    if dy < 0:
        xfrm.set("flipH", "1")
    if dx > 0:
        xfrm.set("flipV", "1")
    xfrm.find(qn("a:off")).set("x", str(int(Inches(cx - w / 2))))
    xfrm.find(qn("a:off")).set("y", str(int(Inches(cy - h / 2))))
    xfrm.find(qn("a:ext")).set("cx", str(int(Inches(w))))
    xfrm.find(qn("a:ext")).set("cy", str(int(Inches(h))))


def glue(c, a, site_a, b, site_b):
    """Record connector endpoints as glued to shapes (so PowerPoint reroutes when a node is dragged)."""
    cnv = c._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
    for tag, shp, idx in (("a:stCxn", a, site_a), ("a:endCxn", b, site_b)):
        el = etree.SubElement(cnv, qn(tag))
        el.set("id", str(shp.shape_id))
        el.set("idx", str(idx))


# ---------------------------------------------------------------- helpers
MARK_RE = re.compile(r"==(.+?)==|\{\{([^{}|]+?)(?:\|([^{}]+))?\}\}")


def parse_marks(t):
    """'a ==b== {{c|d}}' -> [('a ', False, None), ('b', True, None), (' ', False, None), ('d', False, 'c')]"""
    out, pos = [], 0
    for m in MARK_RE.finditer(t):
        if m.start() > pos:
            out.append((t[pos:m.start()], False, None))
        if m.group(1) is not None:
            out.append((m.group(1), True, None))
        else:
            name = m.group(2).strip()
            out.append((m.group(3) or name, False, name))
        pos = m.end()
    if pos < len(t):
        out.append((t[pos:], False, None))
    return out or [("", False, None)]


def plain_text(t):
    return "".join(seg for seg, _, _ in parse_marks(str(t)))


def page_ranges(pages):
    """[2, 3, 4, 7] -> '第 2–4、7 頁'"""
    runs, start = [], None
    for k, pg in enumerate(pages):
        if start is None:
            start = pg
        if k + 1 == len(pages) or pages[k + 1] != pg + 1:
            runs.append(f"{start}–{pg}" if pg > start else str(pg))
            start = None
    return f"第 {'、'.join(runs)} 頁" if runs else ""


def color_table_md(rows):
    lines = ["| 顏色 | 色碼 | 代表 | 出現在 |", "|---|---|---|---|"]
    for r in rows:
        hexes = "／".join(f"`{h}`" for h in r["hex"])
        lines.append(f"| {r['color']} | {hexes} | {r['meaning']} | {page_ranges(r['pages'])} |")
    return "\n".join(lines)


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def label_w(t):
    """Width of a 13pt bold mark label; bold CJK and full-width brackets run ~15% wider than text_w."""
    return text_w(t, 13) * 1.15


def mark_gutters(marks, w, h):
    """Space to keep free beside the image for outside note / band labels."""
    g = {"left": 0.0, "right": 0.0, "top": 0.0, "bottom": 0.0}
    for m in marks:
        side = m.get("side", "right") if m.get("type") == "note" else (
            m.get("label_side", "right") if m.get("type") == "band" and m.get("text") else None)
        if side in ("left", "right"):
            g[side] = min(2.6, w * 0.32)
        elif side in ("top", "bottom"):
            g[side] = min(0.45, h * 0.12)
    return g


def mark_geometry(m):
    """(x, y, w, h) as 0-1 fractions; row/rows and col/cols address equal-height table rows/columns (1-based)."""
    fx, fy, fw, fh = m.get("x"), m.get("y"), m.get("w", 0), m.get("h", 0)
    if "row" in m and m.get("rows"):
        a, b = (m["row"], m["row"]) if isinstance(m["row"], int) else m["row"]
        fy, fh = (a - 1) / m["rows"], (b - a + 1) / m["rows"]
    if "col" in m and m.get("cols"):
        a, b = (m["col"], m["col"]) if isinstance(m["col"], int) else m["col"]
        fx, fw = (a - 1) / m["cols"], (b - a + 1) / m["cols"]
    if m.get("type") == "band" and fx is None:
        fx, fw = 0, 1  # a band spans the full width unless told otherwise
    if fx is None or fy is None:
        return None
    return fx, fy, fw, fh


def set_alpha(sh, pct):
    """Opacity (0-100) of a shape's solid fill."""
    clr = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    for a in clr.findall(qn("a:alpha")):
        clr.remove(a)
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(pct * 1000)))


def fade(sh):
    """Out-of-focus look: fills and lines ~70% transparent, text C0C0C0, pictures washed out."""
    el = sh._element
    if el.tag == qn("p:grpSp"):
        for child in sh.shapes:
            fade(child)
        return
    blip = el.find(".//" + qn("a:blip"))
    if blip is not None:  # picture
        for a in blip.findall(qn("a:alphaModFix")):
            blip.remove(a)
        etree.SubElement(blip, qn("a:alphaModFix")).set("amt", str(FADE_ALPHA * 1000))
    sppr = el.find(qn("p:spPr"))
    if sppr is not None:
        for path in ((qn("a:solidFill"), qn("a:srgbClr")), (qn("a:ln"), qn("a:solidFill"), qn("a:srgbClr"))):
            clr = sppr.find("/".join(path))
            if clr is not None:
                old = clr.find(qn("a:alpha"))
                base = int(old.get("val")) if old is not None else 100000
                if old is not None:
                    clr.remove(old)
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(base * FADE_ALPHA / 100)))
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                r.font.color.rgb = RGBColor.from_string(FADE_TEXT)


def drop_shapes(slide, shape_ids):
    ids = set(shape_ids)
    for sh in list(slide.shapes):
        if sh.shape_id in ids:
            sh._element.getparent().remove(sh._element)


def click_steps(slide, steps):
    """PowerPoint timing tree: on each click the step's "in" shapes fade in and its "out" shapes fade out
    (the first effect is on click, the rest run with it)."""
    def el(parent, tag, **attrs):
        e = etree.SubElement(parent, f"{{{P_NS}}}{tag}")
        for k, v in attrs.items():
            e.set(k, str(v))
        return e

    ids = iter(range(1, 10000))
    timing = el(slide._element, "timing")
    root = el(el(el(timing, "tnLst"), "par"), "cTn", id=next(ids), dur="indefinite", restart="never",
              nodeType="tmRoot")
    seq = el(el(root, "childTnLst"), "seq", concurrent="1", nextAc="seek")
    main = el(seq, "cTn", id=next(ids), dur="indefinite", nodeType="mainSeq")
    clicks = el(main, "childTnLst")
    steps = [st if isinstance(st, dict) else {"in": st, "out": []} for st in steps]
    for stp in steps:
        click = el(el(clicks, "par"), "cTn", id=next(ids), fill="hold")
        el(el(click, "stCondLst"), "cond", delay="indefinite")
        step = el(el(el(click, "childTnLst"), "par"), "cTn", id=next(ids), fill="hold")
        el(el(step, "stCondLst"), "cond", delay="0")
        effects = el(step, "childTnLst")
        k = 0
        for kind, shape_ids in (("out", stp.get("out", [])), ("in", stp.get("in", []))):
            for spid in shape_ids:
                entr = kind == "in"
                eff = el(el(effects, "par"), "cTn", id=next(ids), presetID="10",
                         presetClass="entr" if entr else "exit", presetSubtype="0", fill="hold",
                         grpId="0" if entr else "1", nodeType="clickEffect" if k == 0 else "withEffect")
                k += 1
                el(el(eff, "stCondLst"), "cond", delay="0")
                beh = el(eff, "childTnLst")
                if entr:
                    st = el(beh, "set")
                    cb = el(st, "cBhvr")
                    vis = el(cb, "cTn", id=next(ids), dur="1", fill="hold")
                    el(el(vis, "stCondLst"), "cond", delay="0")
                    el(el(cb, "tgtEl"), "spTgt", spid=spid)
                    el(el(cb, "attrNameLst"), "attrName").text = "style.visibility"
                    el(el(st, "to"), "strVal", val="visible")
                fcb = el(el(beh, "animEffect", transition="in" if entr else "out", filter="fade"), "cBhvr")
                el(fcb, "cTn", id=next(ids), dur="500")
                el(el(fcb, "tgtEl"), "spTgt", spid=spid)
                if not entr:  # hide once the fade-out has finished
                    st = el(beh, "set")
                    cb = el(st, "cBhvr")
                    vis = el(cb, "cTn", id=next(ids), dur="1", fill="hold")
                    el(el(vis, "stCondLst"), "cond", delay="499")
                    el(el(cb, "tgtEl"), "spTgt", spid=spid)
                    el(el(cb, "attrNameLst"), "attrName").text = "style.visibility"
                    el(el(st, "to"), "strVal", val="hidden")
    for tag, evt in (("prevCondLst", "onPrev"), ("nextCondLst", "onNext")):
        el(el(el(seq, tag), "cond", evt=evt, delay="0"), "tgtEl", ).append(etree.Element(f"{{{P_NS}}}sldTgt"))
    # build entries only for text-capable shapes (p:sp), as PowerPoint writes them
    sp_ids = {sh.shape_id for sh in slide.shapes if sh._element.tag == qn("p:sp")}
    bld = el(timing, "bldLst")
    for kind, grp in (("in", "0"), ("out", "1")):
        for stp in steps:
            for spid in stp.get(kind, []):
                if spid in sp_ids:
                    el(bld, "bldP", spid=spid, grpId=grp, animBg="1")


def text_w(t, size):
    em = size / 72
    return sum(em if ord(ch) > 0x2E80 else em * 0.55 for ch in plain_text(t))


def est_height(paras, size, width, space, bullet):
    total = 0
    for p in paras:
        s = p.get("size", size)
        avail = width - (0.3 if bullet else 0)
        lines = 0
        for seg in plain_text(p["t"]).split("\n"):
            lines += max(1, -(-text_w(seg, s) // max(avail, 0.5)))
        total += lines * s * 1.15 / 72 + p.get("space", space) / 72
    return total


def as_list(v):
    if v is None:
        return []
    return v if isinstance(v, list) else [v]


def pick_text(d, key):
    for k in FALLBACK[key]:
        if d.get(k):
            return d[k]
    return None


def norm_fig(f):
    if isinstance(f, dict):
        return f
    f = list(f)
    return {"id": f[0], "caption": f[1] if len(f) > 1 else "", "path": f[2] if len(f) > 2 else None}


def render(spec, base, density, nav, out, keyframes=False, quiet_table=False):
    """Build and save one deck; returns the colour-table rows."""
    WARN.clear()
    deck = Deck(spec, base, density, nav, keyframes)
    prs = deck.build()
    prs.save(out)
    print(f"✔ {out}  ({len(prs.slides)} slides, density={density}, nav={nav})")
    for w in WARN:
        print("  ⚠", w)
    rows = deck.color_table()
    if rows and not quiet_table:
        print("\n色彩對照（貼進回報用）：\n" + color_table_md(rows) + "\n")
    return rows


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("-o", "--out")
    ap.add_argument("--density", choices=list(DENSITIES))
    ap.add_argument("--nav", choices=["single", "double"])
    ap.add_argument("--all-densities", action="store_true")
    ap.add_argument("--keyframes", action="store_true",
                    help="no animations: split each slide with a conclusion bar into before/after slides")
    a = ap.parse_args()
    spec_path = Path(a.spec)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    style = spec.get("style", {})
    nav = a.nav or style.get("nav", "single")
    out = Path(a.out or spec_path.with_suffix(".pptx"))
    if a.all_densities:
        for d in DENSITIES:
            render(spec, spec_path.parent, d, nav, str(out.with_name(f"{out.stem}-{d}.pptx")), a.keyframes,
                   quiet_table=d != list(DENSITIES)[-1])  # print the table once
    else:
        render(spec, spec_path.parent, a.density or style.get("density", "balanced"), nav, str(out),
               a.keyframes)


if __name__ == "__main__":
    sys.exit(main())
